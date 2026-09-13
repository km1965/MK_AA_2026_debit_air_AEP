# -*- coding: utf-8 -*-
"""Export des rapports .docx et .xlsx — MK_A.A 2026.

Nécessite : python-docx, openpyxl (installées via requirements.txt).
"""

import datetime
import os
import tempfile

from . import croquis as croquis_mod
from ..controller import capacite_organe
from ..core.constants import REFERENCES_NORMATIVES, SEUIL_DEPRESSION


def _timestamp() -> str:
    return datetime.datetime.now().strftime("%H:%M:%S")


def _titre_rapport(etat) -> str:
    return "Rapport de synthèse — Débits d'air admis en conduites AEP"


def _info_projet(etat) -> list:
    return [
        ("Maître d'ouvrage", etat.moe),
        ("Projet / marché", etat.projet),
        ("Référence / indice", etat.reference),
        ("Branche(s)", etat.branches),
        ("Schéma de vidange", etat.schema),
        ("Date", datetime.date.today().isoformat() + " " + _timestamp()),
    ]


_NOMS = {"trifon": "Ventouse TRIFON", "ceai": "Clapet SNH",
         "psa": "Purgeur sonique SNH", "vanne": "Vanne sectionnement"}


def _chemin_png_icone() -> str:
    """Chemin vers l'icône du projet (assets/icon.png) — en-tête du .docx.

    Compatible PyInstaller : assets embarqués dans _MEIPASS (onefile) ou
    dossier de l'exécutable (développement : racine du projet).
    """
    import sys
    from pathlib import Path
    if getattr(sys, "frozen", False):
        racines = [Path(getattr(sys, "_MEIPASS", "")),
                   Path(sys.executable).resolve().parent]
    else:
        racines = [Path(__file__).resolve().parent.parent.parent]
    for racine in racines:
        if not racine:
            continue
        png = racine / "assets" / "icon.png"
        if png.exists():
            return str(png)
    return ""


# --- Note MÉTHODOLOGIE MK_A.A (2026) — document Word séparé du rapport ---
_METHODOLOGIE_INTRO = (
    "Le présent paragraphe décrit, à l'intention du contrôleur, la démarche de "
    "calcul appliquée pour l'évaluation des débits d'air à admettre en conduites "
    "AEP lors d'une vidange en casse franche, conformément à la note technique "
    "MK_A.A 2026 (méthode explicite MK_A.A 2026, base NF EN 805).")

_PHASES_METHODOLOGIE = [
    ("Phase 1 — Hypothèses et données d'entrée",
     "Socle de calcul fixe et non modifiable : g = 9,81 m/s², rugosité absolue "
     "k = 0,0005 m, correction de la viscosité cinématique ν en fonction de la "
     "température de l'eau. Seuil de dépression admissible de 3 mCE (NF EN 805) "
     "au point haut de la conduite."),
    ("Phase 2 — Profil en long et géométries",
     "Saisie du profil réel (cotes altimétriques), détermination du diamètre "
     "intérieur D et de la section S = π/4·D², calcul du rapport k/D et des "
     "longueurs des tronçons (a, b, c, d) encadrant chaque point intermédiaire "
     "(PI1 / PI2) et le point haut Ve."),
    ("Phase 3 — Charges H₁ / H₂ aux points intermédiaires (§6.1)",
     "Évaluation des charges H₁ (amont) et H₂ (aval) au regard du seuil de "
     "dépression de 3 mCE. H₁ < 3 → CAS NORMAL (l'eau ne se sépare pas au point "
     "intermédiaire) ; H₁ ≥ 3 → CAS PARTICULIER (formation d'une poche d'air en "
     "PI1, débit complémentaire Q_PI1 à admettre)."),
    ("Phase 4 — Débits par la formule explicite MK_A.A (2026)",
     "Le système de Darcy-Weisbach (ΔH = λ·L/D·V²/2g) couplé à Colebrook-White "
     "(implicite en λ) est résolu sans itération par la formule explicite en "
     "débit : Q = f(ΔH, L, D, ν, k). Application côté amont (Q_am) et/ou côté "
     "aval (Q_av) selon le schéma de vidange choisi (1A/1B, 2A/2B, 3A/3B, 4A/4B). "
     "Le débit d'air admis au point haut est Q_Ve = max(Q_am ; Q_av) pour les "
     "cas combinés, avec la convention de signe Q_PI < 0 → Q_PI = 0 (flux "
     "entrant compensant le drainage)."),
    ("Phase 5 — Vérifications hydrauliques",
     "Contrôle des conditions de validité de la formule : nombre de Reynolds "
     "Re > 4000 et rapport de rugosité k/D ∈ [10⁻⁶ ; 10⁻²], pour chaque tronçon "
     "calculé (Q_am, Q_av, Q_A, Q_B)."),
    ("Phase 6 — Dimensionnement des organes d'air (§8)",
     "Le besoin de référence est le débit d'air requis au point haut, majoré de "
     "la marge de dimensionnement (+15 %), avec plafond d'utilisation de la "
     "capacité installée (90 %) → sur-capacité totale ≈ +28 %. Sélection des "
     "organes par catalogue : "
     "ventouses TRIFON (grand orifice, anti-dépression), clapets d'entrée d'air "
     "CEAI, purgeurs PSA, avec calcul du nombre d'organes en parallèle et du "
     "taux d'utilisation de la capacité installée."),
    ("Phase 7 — Remplissage (§9) et purgeurs soniques SNH",
     "Débit de remplissage de la conduite à V_eau = 2 m/s : Q_remplissage = "
     "V·π/4·D²·3600. Dimensionnement des purgeurs PSA (capacité unitaire "
     "catalogue) et des purgeurs soniques SNH (NSH) selon le modèle physique de "
     "dégazage contrôlé : Q_fill = q_capacité·P_fill/P_svc (détendeur sonique "
     "ancré à la pression de service nominale de 10 bar). L'évacuation massive "
     "de l'air au remplissage est assurée par le grand orifice des ventouses "
     "TRIFON (fonction 3)."),
    ("Phase 8 — Vérification de conformité et cumul casse franche",
     "Vérification de chaque catégorie d'organes (ventouse, clapet, purgeur, "
     "vanne). Contrôle cumulé du groupe d'admission en casse franche : "
     "Q_eau = A·√(2·g·H_z) (débit d'eau de vidange par la brèche), cumul des "
     "débits d'air réels (ventouses + clapets + purgeurs), répartition "
     "amont/aval autour de la vanne de sectionnement (transit par la vanne vs "
     "bypass direct en aval), vitesse de passage de l'air ≤ 40 m/s au point le "
     "plus étroit (anti blocage sonique) et vérification de l'absence "
     "d'étranglement de la vanne (S_vanne > Σ S_amont). Verdict global de "
     "conformité du cumul."),
]

_METHODOLOGIE_NOTE = (
    "toutes les phases reposent sur le socle de calcul fixe (§3) de la note "
    "technique ; la méthode explicite MK_A.A (2026) fournit les débits "
    "sans itération tout en conservant la rigueur du formalisme "
    "Darcy-Weisbach/Colebrook-White (réservé à l'ingénierie hydraulique).")


# ---------------------------------------------------------------------------
# .DOCX
# ---------------------------------------------------------------------------
def _utilise_breche(etat) -> bool:
    """Vrai si la méthode de dimensionnement sélectionnée est la Vidange
    Brèche / Torricelli (le chapitre « casse franche » du rapport n'apparaît
    QUE dans ce cas)."""
    return (getattr(etat, "methode_dimensionnement", "mk_aa") or "mk_aa") == "breche"


def _ecrire_cumul_breche_docx(etat, doc):
    """Chapitre « Cumul des débits d'air — casse franche » (docx).
    Uniquement en méthode de dimensionnement Vidange Brèche/Torricelli."""
    # --- Cumul des débits (note « Débits cumulés ») ---
    doc.add_heading("Cumul des débits d'air — casse franche", level=1)
    doc.add_paragraph(
        "Ce chapitre applique la MÉTHODE VIDANGE BRÈCHE (Torricelli) : on "
        "détermine le débit d'air total à admettre au point haut pour compenser "
        "l'eau vidangée par une rupture totale de la conduite (Q_eau = A·√(2·g·H_z)). "
        "Il est DISTINCT du dimensionnement des ventouses par la formule "
        "MK_A.A (2026) — "
        "qui quantifie le débit d'air admis par un écoulement gravitaire dans "
        "les tronçons. Ici on dimensionne le cumul du GROUPE (ventouses + clapets "
        "+ purgeurs) contre le besoin de la casse franche.")
    cumul = etat.cumul_debits()
    doc.add_paragraph(
        f"Débit d'eau de vidange par la brèche : Q_eau = A·√(2·g·H_z) = "
        f"{cumul['section_conduite']:.3f} × {cumul['v_ecoulement']:.2f} = "
        f"{cumul['besoin_casse']:,.1f} m³/h  (H_z = {cumul['hz']:.1f} m).")
    adm = cumul["admission"]
    doc.add_paragraph(
        f"Cumul admission (ventouses {adm['trifon']:,.0f} + clapets {adm['ceai']:,.0f} "
        f"+ purgeurs {adm['psa']:,.0f}) = {adm['total']:,.0f} m³/h.")
    doc.add_paragraph(
        f"Répartition : amont (transite par la vanne) {adm.get('amont', 0):,.0f} m³/h ; "
        f"aval (bypass conduite) {adm.get('aval', 0):,.0f} m³/h.")
    # Tableau des organes d'admission par position (amont / aval)
    orgs = [(o.get("position") or "amont", _NOMS.get(o.get("type", ""), o.get("type", "")),
             o.get("dn") or "—", o.get("nombre", 1),
             capacit if (capacit := capacite_organe(o, -3)) else "—")
            for o in etat.organes_client]
    if orgs:
        tab = doc.add_table(rows=1, cols=5)
        tab.style = "Light Grid Accent 1"
        for j, htxt in enumerate(["Position", "Type", "DN", "Nbr", "Capacité (m³/h)"]):
            tab.rows[0].cells[j].text = htxt
        for pos, typ, dn, n, cap in orgs:
            cells = tab.add_row().cells
            cells[0].text = str(pos)
            cells[1].text = typ
            cells[2].text = str(dn)
            cells[3].text = str(n)
            cells[4].text = (f"{cap:,.0f}" if isinstance(cap, (int, float)) else str(cap))
    exp = cumul["expulsion"]
    doc.add_paragraph(
        f"Cumul expulsion (remplissage) = {exp['total']:,.0f} m³/h "
        f"(ventouses {exp['trifon']:,.0f} + purgeurs {exp['psa']:,.0f}).")
    vanne = cumul["vanne"]
    if vanne["section_vanne"] > 0:
        sec_txt = (f"Σ sections amont {vanne['section_amont']:.3f} m²"
                   + (f" ; aval {vanne['section_aval']:.3f} m²"
                      if vanne["section_aval"] > 0 else ""))
        doc.add_paragraph(
            f"Vanne sectionnement DN {vanne['dn']:.0f} : S = {vanne['section_vanne']:.3f} m² "
            f"vs {sec_txt} → rapport {vanne['ratio']:.2f} × "
            f"({'sans étranglement' if vanne['conforme'] else 'ÉTRANGLEMENT POTENTIEL'}).")
    vt = cumul["limite_vitesse"]
    v_str = f"{vt['v_air']:.1f} m/s" if vt["v_air"] < 1e6 else "n.c."
    doc.add_paragraph(
        f"Vitesse de passage de l'air : {v_str} à la vanne "
        f"(limite {vt['v_limite']:.0f} m/s — anti blocage sonique)"
        + (f" ; aval {vt['v_aval']:.1f} m/s." if vt.get("v_aval", 0) < 1e6 else "."))
    p = doc.add_paragraph()
    statut = "CONFORME" if (cumul["verdict_admission"] and vanne["conforme"] and vt["ok"]) else "NON CONFORME"
    p.add_run(f"VERDICT CUMUL DES DÉBITS : {statut}").bold = True

    # Montage optimisé proposé lorsque le cumul n'est pas conforme
    if statut != "CONFORME":
        opt = etat.groupe_optimise()
        doc.add_heading("Montage optimisé proposé (TRIFON + SNH)", level=2)
        doc.add_paragraph(
            f"Besoin casse franche : {opt['besoin']:,.0f} m³/h ; admission existante "
            f"{opt['admis_existant']:,.0f} m³/h ; manque à couvrir "
            f"{opt['manque']:,.0f} m³/h.")
        rep = opt["repartition"]
        doc.add_paragraph(
            f"Répartition : amont {rep['amont']:,.0f} m³/h (transite vanne, ≤ "
            f"{rep['kap']:,.0f} m³/h @40 m/s) + aval {rep['aval']:,.0f} m³/h "
            f"(bypass conduite).")
        for o in opt["groupe"]:
            posf = f" [{o['position']}]" if o.get("position") else ""
            doc.add_paragraph(
                f"• {o['nombre']} × {_NOMS.get(o['type'], o['type'])} "
                f"DN{o['dn']}{posf}", style="List Bullet")
        if opt["vanne_dn"]:
            doc.add_paragraph(
                f"Vanne de sectionnement recommandée : DN {opt['vanne_dn']} "
                f"(= DN de la conduite principale, composant neutre).")
        doc.add_paragraph(
            f"Couverture totale {opt['coverage']:,.0f} m³/h → "
            f"{'CONFORME (vitesses ≤ 40 m/s)' if opt['conforme'] else 'à compléter'}.")
        if opt["note"]:
            doc.add_paragraph(opt["note"])

    # Références normatives — fin du chapitre casse franche
    doc.add_heading("Références normatives — Admission d'air (casse franche)", level=2)
    for ref_text in [
        "NF EN 805 (2000) — Prévention des dépressions dans les réseaux AEP, "
        "dépression admissible ≤ 3 mCE.",
        "AWWA M51 (2020) — Air release and vacuum valves for waterworks.",
        "Torricelli — Formule Q = A·√(2·g·H) pour le débit de vidange par "
        "brèche (écoulement d'un orifice sous charge).",
        "Critère anti-étranglement sonique : v_air < 40 m/s (AWWA).",
    ]:
        doc.add_paragraph(ref_text, style="List Bullet")


def _ecrire_debits_cumules_docx(etat, doc):
    """Chapitre « Débits cumulés » (docx) — méthode MK_A.A.

    Calqué sur la note « Débits cumulés » de référence : la seule différence
    est que le débit d'air requis (formule Q_air = A·√(2·g·H_z) de la brèche)
    est NÉGLIGÉ et remplacé par le débit Q MK_A.A calculé (Q_Ve).
    """
    try:
        ga = etat.verifier_groupe_mk_aa()
        cum = etat.cumul_debits()
    except Exception as e:
        doc.add_heading("Débits cumulés", level=1)
        doc.add_paragraph(f"Vérification des débits cumulés indisponible : {e}")
        return
    dc = cum["dc"]; adm = ga["admission"]; exp = cum["expulsion"]
    vanne = cum["vanne"]; vt = cum["limite_vitesse"]
    besoin_tot = ((etat.resultat.q_ve_m3h or 0.0)
                  + (etat.resultat.q_pi1_m3h or 0.0)
                  + (etat.resultat.q_pi2_m3h or 0.0))

    doc.add_heading("Débits cumulés — vérification du groupe d'air au point haut",
                    level=1)
    doc.add_paragraph(
        "Pour vérifier la capacité d'admission ou d'échappement d'air à un point "
        "haut, on peut cumuler les débits d'air individuels de chaque appareil, "
        "mais sous trois conditions hydrauliques et géométriques majeures. Le "
        "dimensionnement s'effectue sur la somme des débits d'air réels admis ou "
        "expulsés, et non sur la simple somme des diamètres des appareils.")

    # 1. Règle du cumul
    doc.add_heading("Règle du cumul des débits d'air", level=2)
    doc.add_paragraph(
        "Le débit total à l'admission (anti-dépression lors d'une vidange rapide "
        "ou d'un coup de bélier) ou à l'expulsion (remplissage de la conduite) "
        "est la somme directe des capacités de chaque composant à la même "
        f"différence de pression (−{SEUIL_DEPRESSION:g} mCE) :")
    doc.add_paragraph("Pour l'admission d'air (protection contre la dépression) : "
                      "les clapets d'admission d'air et les ventouses/purgeurs "
                      "(en phase d'entrée d'air) travaillent en parallèle.",
                      style="List Bullet")
    doc.add_paragraph("Pour l'expulsion d'air (remplissage / régime permanent) : "
                      "seuls les éléments conçus pour l'échappement d'air "
                      "(ventouses et purgeurs soniques) interviennent.",
                      style="List Bullet")
    doc.add_paragraph(
        f"ΣQ admis = ventouses {adm['trifon']:,.0f} + clapets {adm['ceai']:,.0f} "
        f"+ purgeurs {adm['psa']:,.0f} = {adm['total']:,.0f} m³/h.")
    doc.add_paragraph(
        f"ΣQ expulsé = ventouses {exp['trifon']:,.0f} + purgeurs {exp['psa']:,.0f} "
        f"= {exp['total']:,.0f} m³/h.")

    # 2. Condition essentielle : tubulure + vanne
    doc.add_heading("Condition essentielle : la tubulure de raccordement et la "
                    "vanne de sectionnement", level=2)
    doc.add_paragraph(
        "Le point critique réside dans la vanne de sectionnement sous l'ensemble : "
        "la section de passage minimale du collecteur (la vanne et le piquage "
        "sur le DN principal) ne doit pas créer une perte de charge supérieure à "
        "la capacité globale des appareils installés au-dessus.")
    doc.add_paragraph(
        f"Section de la vanne DN {vanne['dn']:g} = {vanne['section_vanne']:.3f} m². "
        f"La somme des sections des appareils en partie haute (DN réels + col "
        f"sonique des purgeurs) représente environ {vanne['section_amont']:.3f} m². "
        f"S_vanne/ΣS = {vanne['ratio']:.2f} → "
        + ("la vanne est largement dimensionnée et ne bridera pas le passage "
           "d'air vers les équipements." if vanne["conforme"]
           else "ÉTRANGLEMENT POTENTIEL — la vanne bridera le passage d'air."))

    # 3. Points de vigilance
    doc.add_heading("Points de vigilance pour la vérification de calcul", level=2)
    doc.add_paragraph("Principe hydraulique (vidange / coup de bélier / rupture "
                      "brutale en aval du point haut) : la lame d'eau s'échappe "
                      "par gravité à grande vitesse, créant un effet piston "
                      "immédiatement sous le dôme. Les appareils du bloc au-dessus "
                      "de la vanne s'ouvrent en grand sous l'effet de la dépression "
                      "instantanée et aspirent l'air en parallèle ; le débit d'air "
                      "global admis est la somme exacte des débits d'air "
                      "individuels réels de chaque appareil. Le cumul est "
                      "parfaitement valable et constitue la base même de la "
                      "protection contre l'écrasement (flambement) de la conduite.",
                      style="List Number")
    doc.add_paragraph(
        "Débit d'air requis pour la vérification : Q MK_A.A calculé. "
        "(La formule du débit d'air d'entrée requis par la brèche — "
        "Q_air = A·√(2·g·H_z), Torricelli — est NÉGLIGÉE pour la vérification ; "
        "la référence retenue est le débit MK_A.A Q_Ve, débit d'air admis par "
        "l'écoulement gravitaire des tronçons au point haut, formule MK_A.A "
        "2026.) "
        f"Q_Ve = {etat.resultat.q_ve_m3h:,.1f} m³/h · Q_PI1 = "
        f"{etat.resultat.q_pi1_m3h:,.1f} m³/h · Q_PI2 = "
        f"{etat.resultat.q_pi2_m3h:,.1f} m³/h → besoin réel du tronçon "
        f"(casse franche, SANS majoration) = {besoin_tot:,.1f} m³/h.",
        style="List Number")
    v_text = (f"{vt['v_air']:.1f} m/s" if vt["v_air"] < 1e6 else "n.c.")
    doc.add_paragraph(
        "Vérification des conditions limites : en premier lieu, extraire sur les "
        f"courbes constructeurs le débit d'air de chaque appareil pour la "
        f"dépression de référence (−{SEUIL_DEPRESSION:g} mCE) sans dépasser cette "
        f"valeur pour éviter l'instabilité de la conduite ou le blocage sonique ; "
        f"en second lieu, la contrainte de vitesse au niveau de la "
        f"tubulure/vanne : v = {v_text} ≤ {vt['v_limite']:.0f} m/s "
        + ("→ aucune perte de charge limitante (régime sous-sonique)."
           if vt["ok"] else "→ RÉGIME SONIQUE, débit réel réduit."),
        style="List Number")

    # 4. Bilan
    doc.add_heading("Bilan de vérification", level=2)
    fs_tot = (adm["total"] / besoin_tot
              if besoin_tot > 0 else (float("inf") if adm["total"] > 0 else 0.0))
    verdict = (fs_tot >= 1.0) and vanne["conforme"] and vt["ok"]
    p = doc.add_paragraph()
    p.add_run(f"VERDICT DÉBITS CUMULÉS : {('CONFORME' if verdict else 'NON CONFORME')}"
              ).bold = True
    doc.add_paragraph(
        f"• Groupe d'admission {adm['total']:,.0f} m³/h vs besoin réel du tronçon "
        f"{besoin_tot:,.1f} m³/h (Q MK_A.A, sans majoration) → "
        f"Fs = {fs_tot:.2f} → "
        f"{'COUVERT (Fs ≥ 1,00 — le contrôleur juge la valeur du Fs)' if fs_tot >= 1.0 else 'INSUFFISANT (Fs < 1,00)'}.")
    doc.add_paragraph(
        f"• Vanne/sections : S vanne {vanne['section_vanne']:.3f} m² vs ΣS appareils "
        f"{vanne['section_amont']:.3f} m² → "
        f"{'sans étranglement' if vanne['conforme'] else 'étranglement'}.")
    doc.add_paragraph(
        f"• Vitesse d'air au point le plus étroit : {v_text} ≤ {vt['v_limite']:.0f} "
        f"m/s → {'sous-sonique OK' if vt['ok'] else 'blocage sonique'}.")
    doc.add_heading("Références normatives — Débits cumulés", level=2)
    for ref_text in [
        "NF EN 805 (2000) — Prévention des dépressions dans les réseaux AEP, "
        f"dépression admissible ≤ {SEUIL_DEPRESSION:g} mCE.",
        "AWWA M51 (2020) — Air release and vacuum valves for waterworks.",
        "MK_A.A (2026) — Débit d'air admis par l'écoulement gravitaire "
        "des tronçons (Q_Ve) retenu pour la vérification du cumul.",
        "Critère anti-étranglement sonique : v_air < 40 m/s (AWWA).",
    ]:
        doc.add_paragraph(ref_text, style="List Bullet")


def _ecrire_cumul_breche_xlsx(etat, ws, wb, titre_font, ok_fill, nok_fill, autowidth):
    """Feuille « Cumul débits » (casse franche) — seulement en méthode Brèche."""
    from openpyxl.styles import Font
    # --- Feuille Cumul des débits (casse franche) ---
    ws = wb.create_sheet("Cumul débits")
    ws["A1"] = "Cumul des débits d'air — casse franche"; ws["A1"].font = titre_font
    cumul = etat.cumul_debits()
    adm = cumul["admission"]; exp = cumul["expulsion"]; vanne = cumul["vanne"]
    vt = cumul["limite_vitesse"]
    rows = [
        ("Dénivelé H_z (m)", cumul["hz"]),
        ("Section conduite (m²)", round(cumul["section_conduite"], 4)),
        ("Vitesse d'écoulement (m/s)", round(cumul["v_ecoulement"], 2)),
        ("Débit d'eau de vidange Q_eau (m³/h)", round(cumul["besoin_casse"], 1)),
        ("Cumul admission — ventouses (m³/h)", round(adm["trifon"], 1)),
        ("Cumul admission — clapets (m³/h)", round(adm["ceai"], 1)),
        ("Cumul admission — purgeurs (m³/h)", round(adm["psa"], 1)),
        ("TOTAL admission (m³/h)", round(adm["total"], 1)),
        ("Admission amont (transite vanne, m³/h)", round(adm.get("amont", 0), 1)),
        ("Admission aval (bypass conduite, m³/h)", round(adm.get("aval", 0), 1)),
        ("TOTAL expulsion / remplissage (m³/h)", round(exp["total"], 1)),
        ("DN vanne sectionnement (mm)", vanne["dn"] or "—"),
        ("Section vanne (m²)", round(vanne["section_vanne"], 3) if vanne["section_vanne"] else "—"),
        ("Σ sections organes amont (m²)", round(vanne["section_amont"], 3)),
        ("Rapport vanne / Σ sections amont", round(vanne["ratio"], 2) if vanne["ratio"] != float("inf") else "—"),
        ("Capacité de passage à 40 m/s (m³/h)", round(vt["capacite_passage_40"], 1)),
        ("Vitesse air à la vanne (m/s)", round(vt["v_air"], 1) if vt["v_air"] < 1e6 else "∞"),
        ("Vitesse air aval (m/s)", round(vt.get("v_aval", 0), 1) if vt.get("v_aval", 0) < 1e6 else "∞"),
        ("Limite vitesse (m/s)", vt["v_limite"]),
    ]
    for i, (lib, val) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=lib).font = Font(bold=True)
        ws.cell(row=i, column=2, value=val)
    verdict = cumul["verdict_admission"] and vanne["conforme"] and vt["ok"]
    vc = ws.cell(row=3 + len(rows) + 1, column=1,
                 value="VERDICT CUMUL DES DÉBITS")
    vc.font = Font(bold=True, size=12)
    vcell = ws.cell(row=3 + len(rows) + 1, column=2,
                    value="CONFORME" if verdict else "NON CONFORME")
    vcell.font = Font(bold=True, color="FFFFFF")
    vcell.fill = ok_fill if verdict else nok_fill
    if not verdict:
        opt = etat.groupe_optimise()
        base = 3 + len(rows) + 3
        ws.cell(row=base, column=1, value="MONTAGE OPTIMISÉ PROPOSÉ (TRIFON + SNH)").font = Font(bold=True, size=12)
        ws.cell(row=base + 1, column=1, value="Besoin / admit existant / manque (m³/h)")
        ws.cell(row=base + 1, column=2, value=f"{opt['besoin']:,.0f} / {opt['admis_existant']:,.0f} / {opt['manque']:,.0f}")
        ws.cell(row=base + 2, column=1, value="Répartition amont / aval (m³/h)")
        wsp = ws.cell(row=base + 2, column=2,
                      value=f"{opt['repartition']['amont']:,.0f} / {opt['repartition']['aval']:,.0f}")
        ws.cell(row=base + 3, column=1, value="Vanne sectionnement recommandée")
        ws.cell(row=base + 3, column=2,
                value=f"DN {opt['vanne_dn']}" if opt.get("vanne_dn") else "—")
        j = base + 4
        for o in opt["groupe"]:
            posf = f" [{o['position']}]" if o.get("position") else ""
            ws.cell(row=j, column=1,
                    value=f"{o['nombre']} × {_NOMS.get(o['type'], o['type'])} DN{o['dn']}{posf}")
            j += 1
        ws.cell(row=j, column=1, value="Couverture totale (m³/h)").font = Font(bold=True)
        ws.cell(row=j, column=2, value=round(opt["coverage"], 1))
        ws.cell(row=j + 1, column=1, value="Montage").font = Font(bold=True)
        ws.cell(row=j + 1, column=2, value="CONFORME" if opt["conforme"] else "À COMPLÉTER")
    # Références normatives — fin feuille cumul (après verdict + montage éventuel)
    ref_base = (j + 2) if not verdict else (3 + len(rows) + 2)
    ws.cell(row=ref_base, column=1,
            value="Références normatives — Admission d'air (casse franche)").font = Font(bold=True, size=11)
    for k, ref in enumerate([
        "NF EN 805 (2000) — Prévention des dépressions ≤ 3 mCE.",
        "AWWA M51 (2020) — Air release and vacuum valves.",
        "Torricelli — Q = A·√(2·g·H).",
        "Critère anti-étranglement sonique : v < 40 m/s.",
    ], start=1):
        ws.cell(row=ref_base + k, column=1, value=ref)
    autowidth(ws, 2)


def _ecrire_debits_cumules_xlsx(etat, ws, wb, titre_font, ok_fill, nok_fill, autowidth):
    """Feuille « Débits cumulés » (méthode MK_A.A — Q MK_A.A pour la
    vérification, au lieu de la formule Q_air de la brèche)."""
    from openpyxl.styles import Font
    ws = wb.create_sheet("Débits cumulés")
    ws["A1"] = "Débits cumulés — vérification du groupe d'air au point haut"
    ws["A1"].font = titre_font
    ga = etat.verifier_groupe_mk_aa()
    cumul = etat.cumul_debits()
    adm = ga["admission"]; exp = cumul["expulsion"]
    vanne = cumul["vanne"]; vt = cumul["limite_vitesse"]
    rows = [
        ("Q MK_A.A requis Q_Ve (m³/h)", round(ga["besoin_mk_aa"], 1)),
        ("Q_PI1 / Q_PI2 (m³/h)",
         f"{etat.resultat.q_pi1_m3h:,.1f} / {etat.resultat.q_pi2_m3h:,.1f}"),
        ("Besoin réel du tronçon — sans majoration (m³/h)",
         round(etat.resultat.q_ve_m3h + etat.resultat.q_pi1_m3h
               + etat.resultat.q_pi2_m3h, 1)),
        ("Fs du groupe d'admission (Capacité / Besoin réel)",
         round(adm["total"] / (etat.resultat.q_ve_m3h + etat.resultat.q_pi1_m3h
                               + etat.resultat.q_pi2_m3h), 2)
         if (etat.resultat.q_ve_m3h + etat.resultat.q_pi1_m3h
             + etat.resultat.q_pi2_m3h) > 0 else "—"),
        ("Cumul admission — ventouses (m³/h)", round(adm["trifon"], 1)),
        ("Cumul admission — clapets (m³/h)", round(adm["ceai"], 1)),
        ("Cumul admission — purgeurs (m³/h)", round(adm["psa"], 1)),
        ("TOTAL admission (m³/h)", round(adm["total"], 1)),
        ("Admission amont (transite vanne, m³/h)", round(ga["amont"], 1)),
        ("Admission aval (bypass conduite, m³/h)", round(ga["aval"], 1)),
        ("TOTAL expulsion / remplissage (m³/h)", round(exp["total"], 1)),
        ("DN vanne sectionnement (mm)", vanne["dn"] or "—"),
        ("Section vanne (m²)", round(vanne["section_vanne"], 3) if vanne["section_vanne"] else "—"),
        ("Σ sections organes en partie haute (m²)", round(vanne["section_amont"], 3)),
        ("Rapport vanne / Σ sections appareils",
         round(vanne["ratio"], 2) if vanne["ratio"] != float("inf") else "—"),
        ("Capacité de passage à 40 m/s (m³/h)", round(vt["capacite_passage_40"], 1)),
        ("Vitesse air vanne / aval (m/s)",
         f"{round(vt['v_air'], 1) if vt['v_air'] < 1e6 else '∞'} / "
         f"{round(vt.get('v_aval', 0), 1) if vt.get('v_aval', 0) < 1e6 else '∞'}"),
        ("Limite vitesse anti blocage sonique (m/s)", vt["v_limite"]),
    ]
    for i, (lib, val) in enumerate(rows, start=3):
        ws.cell(row=i, column=1, value=lib).font = Font(bold=True)
        ws.cell(row=i, column=2, value=val)
    bes_tot = (etat.resultat.q_ve_m3h + etat.resultat.q_pi1_m3h
               + etat.resultat.q_pi2_m3h)
    fs_tot = (adm["total"] / bes_tot if bes_tot > 0
              else (float("inf") if adm["total"] > 0 else 0.0))
    verdict = (fs_tot >= 1.0) and vanne["conforme"] and vt["ok"]
    vcell = ws.cell(row=3 + len(rows) + 1, column=2,
                    value="CONFORME" if verdict else "NON CONFORME")
    ws.cell(row=3 + len(rows) + 1, column=1,
            value="VERDICT DÉBITS CUMULÉS").font = Font(bold=True, size=12)
    vcell.font = Font(bold=True, color="FFFFFF")
    vcell.fill = ok_fill if verdict else nok_fill
    ref_base = 3 + len(rows) + 3
    ws.cell(row=ref_base, column=1,
            value="Références normatives — Débits cumulés").font = Font(bold=True, size=11)
    for k, ref in enumerate([
        "NF EN 805 (2000) — Dépressions ≤ 3 mCE.",
        "AWWA M51 (2020) — Air release and vacuum valves.",
        "MK_A.A (2026) — Q_Ve retenu pour la vérification du cumul "
        "(formule Q_air / Torricelli négligée).",
        "Critère anti-étranglement sonique : v < 40 m/s.",
    ], start=1):
        ws.cell(row=ref_base + k, column=1, value=ref)
    autowidth(ws, 2)


def exporter_docx(etat, chemin: str) -> str:
    """Génère un rapport .docx complet (texte + croquis image + tableaux)."""
    etat.calculer()  # le rapport doit toujours refléter l'état courant
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Styles de base
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # --- En-tête : icône du projet + référence ---
    try:
        sec = doc.sections[0]
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        icone_png = _chemin_png_icone()
        if icone_png:
            run_img = hp.add_run()
            try:
                from PIL import Image as _PILImage
                tmp_ico = os.path.join(tempfile.gettempdir(),
                                       "_mk_aa_docx_icone_petite.png")
                with _PILImage.open(icone_png) as im:
                    im.convert("RGBA").resize((64, 64), _PILImage.LANCZOS).save(tmp_ico)
                run_img.add_picture(tmp_ico, height=Inches(0.38))
            except Exception:
                run_img.add_picture(icone_png, height=Inches(0.38))
            run_txt = hp.add_run("   MK_A.A 2026 — Débits d'air admis en conduites AEP")
        else:
            run_txt = hp.add_run("MK_A.A 2026 — Débits d'air admis en conduites AEP")
        run_txt.font.size = Pt(9)
        run_txt.font.color.rgb = RGBColor(0x1F, 0x6F, 0xEB)
    except Exception as ex:  # l'en-tête ne doit jamais bloquer l'export
        doc.add_paragraph(f"(En-tête indisponible : {ex})")

    # --- Titre ---
    h = doc.add_heading(_titre_rapport(etat), level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Note technique MK_A.A 2026 — Casse franche").alignment = \
        WD_ALIGN_PARAGRAPH.CENTER

    # --- Identification ---
    doc.add_heading("Identification du projet", level=1)
    for lib, val in _info_projet(etat):
        p = doc.add_paragraph()
        p.add_run(f"{lib} : ").bold = True
        p.add_run(str(val))

    # --- Tableau récapitulatif des entrées (§1) ---
    doc.add_heading("Tableau récapitulatif des entrées", level=1)
    try:
        ent = etat.tableau_entrees()
        t_ent = doc.add_table(rows=1, cols=3)
        t_ent.style = "Light Grid Accent 1"
        for j, h in enumerate(["Paramètre", "Valeur", "Unité"]):
            t_ent.rows[0].cells[j].text = h
        params = [
            ("Diamètre nominal DN", f"{ent['dn_mm']}", "mm"),
            ("Diamètre intérieur D", f"{ent['d_m']:.4f}" if ent['d_m'] else "—", "m"),
            ("Section S = π/4·D²", f"{ent['section_m2']:.4f}" if ent['section_m2'] else "—", "m²"),
            ("Rugosité k", f"{ent['k_m']}", "m"),
            ("Rapport k/D", f"{ent['kd']:.3e}" if ent['kd'] else "—", "—"),
            ("Schéma de vidange", ent['schema'], "—"),
            ("Altitude Ve (point haut)", f"{ent['z_ve']}", "m NGM"),
            ("Altitude Vi1", f"{ent['z_vi1']}" if ent['z_vi1'] else "—", "m NGM"),
            ("Altitude Vi2", f"{ent['z_vi2']}" if ent['z_vi2'] else "—", "m NGM"),
            ("ΔH casse franche (hz)", f"{ent['hz_casse_franche']}" if ent['hz_casse_franche'] else "—", "m"),
            ("DN vanne de sectionnement", f"{ent['dn_vanne']}" if ent['dn_vanne'] else "—", "mm"),
            ("PN pression nominale conduite", f"{ent['pression_nominale']}", "bar"),
            ("Température eau", f"{ent['temperature']}", "°C"),
            ("Viscosité cinématique ν", f"{ent['nu']:.3e}" if ent['nu'] else "—", "m²/s"),
            ("Dépression admissible ΔH", f"{ent['delta_h']}", "mCE"),
        ]
        params += [(libd, (f"{vald}" if vald else "—"), unitd)
                   for libd, vald, unitd in ent.get("distances", [])]
        for lib, val, unit in params:
            r = t_ent.add_row().cells
            r[0].text = lib; r[1].text = str(val); r[2].text = unit
        if ent["organes"]:
            p = doc.add_paragraph()
            p.add_run("Organes proposés par le client :").bold = True
            t_org = doc.add_table(rows=1, cols=5)
            t_org.style = "Light Grid Accent 1"
            for j, h in enumerate(["Type", "DN", "Qté", "Fournisseur", "Position"]):
                t_org.rows[0].cells[j].text = h
            for o in ent["organes"]:
                r = t_org.add_row().cells
                r[0].text = o['type']; r[1].text = f"DN{o['dn']}"
                r[2].text = str(o['nombre']); r[3].text = o['fournisseur']
                r[4].text = o['position']
    except Exception as e:
        doc.add_paragraph(f"(Tableau des entrées indisponible : {e})")

    # --- Références normatives (complètes) ---
    doc.add_heading("Références normatives", level=1)
    doc.add_paragraph("Les calculs et la vérification s'appuient sur les textes suivants :")
    try:
        refs = etat.references_completes()
    except Exception:
        refs = list(REFERENCES_NORMATIVES)
    for ref in refs:
        p = doc.add_paragraph(ref, style="List Bullet")

    # --- Hypothèses ---
    doc.add_heading("Hypothèses de calcul et profil en long", level=1)
    dc = etat.dc()
    nu = etat.derniere_trace["nu"] if etat.derniere_trace else 1.02e-6
    doc.add_paragraph("Socle fixe : g = 9,81 m/s² · k = 0,0005 m · "
                      "seuil de dépression = 3 mCE (NF EN 805)")
    doc.add_paragraph(f"Température = {etat.temperature:.1f} °C → ν = {nu:.3e} m²/s")
    doc.add_paragraph("Méthode : Formule explicite MK_A.A (2026) — validité "
                      "Re > 4000 et k/D ∈ [10⁻⁶ ; 10⁻²].")

    # --- Croquis du profil (image) ---
    doc.add_heading("Croquis du profil en long", level=2)
    try:
        tmp = os.path.join(tempfile.gettempdir(), "_mk_aa_docx_croquis.png")
        croquis_mod.croquis_png(etat, tmp)
        doc.add_picture(tmp, width=Inches(6.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    except Exception as ex:
        doc.add_paragraph(f"(Croquis indisponible : {ex})")

    # --- Conduite ---
    doc.add_heading("Caractéristiques de la conduite", level=1)
    t = doc.add_table(rows=1, cols=2)
    t.style = "Light Grid Accent 1"
    hdr = t.rows[0].cells
    hdr[0].text = "Paramètre"; hdr[1].text = "Valeur"
    for g, vals in [("DN", f"{dc.dn:g} mm"), ("D intérieur", f"{dc.d:.4f} m"),
                    ("Section S", f"{dc.section:.4f} m²"), ("k/D", f"{dc.kd:.3e}")]:
        r = t.add_row().cells
        r[0].text = g
        r[1].text = vals

    # --- Journal de calcul ---
    doc.add_heading("Calcul — journal pas-à-pas", level=1)
    res = etat.resultat
    for step in res.journal:
        doc.add_paragraph(step, style="List Bullet")
    p = doc.add_paragraph()
    p.add_run(f"RÉSULTAT : Q_Ve = {res.q_ve_m3h:,.1f} m³/h "
              f"(amont {res.q_am_m3h:,.1f} · aval {res.q_av_m3h:,.1f} "
              f"· PI1 {res.q_pi1_m3h:,.1f} · PI2 {res.q_pi2_m3h:,.1f})").bold = True

    # --- Remplissage ---
    doc.add_heading("Remplissage (§9)", level=1)
    remp = etat.derniere_trace["remplissage"]
    doc.add_paragraph(f"Débit de remplissage (V = 2 m/s) : {remp['q']:,.0f} m³/h")
    p = remp["purgeurs"]
    doc.add_paragraph(f"Purgeur PSA : {p.nombre} × {p.nom} DN{p.dn} "
                      f"(cap. unit. {p.capacite_unitaire_m3h:,.0f} m³/h)")
    sn = remp.get("purgeurs_snh")
    if sn is not None:
        doc.add_paragraph(
            f"Purgeur sonique SNH (NSH) — dimensionnement physique au remplissage : "
            f"{sn.nombre} × {sn.nom} DN{sn.dn} "
            f"(Q_fill sonique unit. {sn.capacite_unitaire_m3h:,.1f} m³/h, "
            f"modèle Q_fill = q_cap·(P_fill/P_svc)). "
            f"Le débit physique de ces purgeurs à P≈atm est faible : l'évacuation "
            f"massive au remplissage revient au grand orifice des ventouses TRIFON.")

    # --- Implantation & justification amont/aval ---
    doc.add_heading("Implantation & justification amont/aval", level=1)
    try:
        plan = etat.plan_pose()
        doc.add_paragraph(plan["repere"])
        pose = doc.add_table(rows=1, cols=2)
        pose.style = "Light Grid Accent 1"
        pose.rows[0].cells[0].text = "Position (par rapport à la vanne de sectionnement)"
        pose.rows[0].cells[1].text = "Organes"
        for pos, lib in (("amont", "AMONT — côté point haut Ve"),
                         ("aval", "AVAL — côté conduite"),
                         ("pi1", "PI1 — point intermédiaire amont (organes dédiés)"),
                         ("pi2", "PI2 — point intermédiaire aval (organes dédiés)"),
                         ("sans_position", "Sans position")):
            items = plan["organes_par_position"].get(pos, [])
            r = pose.add_row().cells
            r[0].text = lib
            r[1].text = (" ; ".join(items) if items else "—")
        if plan.get("combo"):
            for j in plan["combo"].get("justifications", []):
                doc.add_paragraph(j, style="List Bullet")
    except Exception as e:
        doc.add_paragraph(f"(Implantation indisponible : {e})")

    # --- Plan d'implantation X, Y, Z (exécution) ---
    doc.add_heading("Plan d'implantation — coordonnées X, Y, Z (exécution)", level=1)
    try:
        xyz = etat.plan_xyz()
        doc.add_paragraph(xyz["base"])
        t = doc.add_table(rows=1, cols=8)
        t.style = "Light Grid Accent 1"
        for j, h in enumerate(["Rép.", "Organe", "DN", "Nbr", "Position",
                               "X (m)", "Y (m)", "Z (m)"]):
            t.rows[0].cells[j].text = h
        for o in xyz["organes"]:
            r = t.add_row().cells
            r[0].text = o["ref"]
            r[1].text = o["type"]
            r[2].text = f"DN{o['dn']}"
            r[3].text = str(o["nombre"])
            r[4].text = o["position"]
            r[5].text = f"{o['x_m']:.1f}"
            r[6].text = f"{o['y_m']:.1f}"
            r[7].text = (f"{o['z_m']:.2f}" if o["z_m"] is not None else "—")
        # Remarques (déportées sous le tableau pour lisibilité)
        for o in xyz["organes"]:
            p = doc.add_paragraph()
            p.add_run(f"{o['ref']} — {o['remarque']}").italic = True
    except Exception as e:
        doc.add_paragraph(f"(Plan d'implantation indisponible : {e})")

    # --- Vérification organes client ---
    doc.add_heading("Vérification des organes (proposition client)", level=1)
    res = etat.resultat
    q_p = [("Q_Ve", res.q_ve_m3h or 0.0),
           ("Q_PI1", res.q_pi1_m3h or 0.0),
           ("Q_PI2", res.q_pi2_m3h or 0.0)]
    q_p = [(l, v) for l, v in q_p if v > 0]
    besoin_tot = sum(v for _, v in q_p)
    doc.add_paragraph(
        "Vérification de la proposition client sur le BESOIN BRUT en casse "
        "franche (débits MK_A.A 2026, SANS majoration) : Fs = Capacité "
        "installée / Demande du point. Fs ≥ 1,00 → catégorie conforme ; le "
        "contrôleur juge la valeur du Fs — une insuffisance (Fs < 1,00) ou une "
        "exigence de Fs supérieur est portée dans ses écrits et observations. "
        "Les purgeurs soniques (évacuation au remplissage) n'ont pas de Fs.")
    doc.add_paragraph(
        "Besoins bruts par point : "
        + (" ; ".join(f"{l} = {v:,.2f} m³/h" for l, v in q_p) or "—")
        + f" → besoin réel du tronçon = {besoin_tot:,.2f} m³/h (casse franche).")
    ver = doc.add_table(rows=1, cols=10)
    ver.style = "Light Grid Accent 1"
    for j, htxt in enumerate(["Repére", "Demande (m³/h)", "Catégorie", "DN (mm)",
                              "Nbr", "Implantation", "Fournisseur",
                              "Capacité (m³/h)", "Fs", "Verdict"]):
        ver.rows[0].cells[j].text = htxt
    for r in etat.tableau_verification_organes():
        c = ver.add_row().cells
        c[0].text = r["repere"]
        c[1].text = (f"{r['demande_m3h']:.2f}" if r["demande_m3h"] is not None else "")
        c[2].text = r["categorie"]
        c[3].text = str(r["dn"] if r["dn"] is not None else "—")
        c[4].text = str(r["nombre"])
        c[5].text = r["implantation"]
        c[6].text = r["fournisseur"]
        c[7].text = (f"{r['capacite_m3h']:,.2f}" if r["capacite_m3h"] is not None else "—")
        c[8].text = (f"{r['fs']:.2f}" if r["fs"] is not None else "")
        c[9].text = r["verdict"]

    # --- Groupe retenu ---
    try:
        comp = etat.proposition_client_justifiee()
    except Exception:
        comp = None
    if comp is not None:
        doc.add_heading("Dimensionnement corrigé — Groupe retenu", level=2)
        doc.add_paragraph(comp["verdict_global"])
        t2 = doc.add_table(rows=1, cols=6)
        t2.style = "Light Grid Accent 1"
        for j, htxt in enumerate(["Rép.", "Catégorie", "DN (mm)", "Nbr",
                                  "Fournisseur", "Rôle / capacité"]):
            t2.rows[0].cells[j].text = htxt
        for o in comp["retenu"]:
            cells = t2.add_row().cells
            cells[0].text = o["rep"]
            cells[1].text = o["categorie"]
            cells[2].text = str(o["dn"] if o["dn"] is not None else "—")
            cells[3].text = str(o["nombre"])
            cells[4].text = o["fournisseur"] or "—"
            cells[5].text = o["role"]

    # --- Cumul des débits : casse franche (Brèche) OU débits cumulés Q MK_A.A ---
    if _utilise_breche(etat):
        _ecrire_cumul_breche_docx(etat, doc)
    else:
        _ecrire_debits_cumules_docx(etat, doc)
    # --- Vérifications structurelles et aérauliques (§4) ---
    doc.add_heading("Vérifications structurelles et aérauliques", level=1)

    # 4.1 Flambement (collapse)
    doc.add_heading("Résistance au vide — Flambement", level=2)
    doc.add_paragraph(
        "La dépression maximale admissible dans la conduite ne doit pas provoquer "
        "le collapse (flambement) du tube sous l'effet de la pression externe ou "
        "du vide intérieur. Critère simplifié (AWWA M11 Ch.6 / NF EN 1295-1) : "
        "ΔH_max < 50% × P_nominale (convertie en mCE).")
    try:
        fl = etat.verifier_flambement()
        t_fl = doc.add_table(rows=1, cols=2)
        t_fl.style = "Light Grid Accent 1"
        t_fl.rows[0].cells[0].text = "Paramètre"; t_fl.rows[0].cells[1].text = "Valeur"
        for lib, val in [
            ("ΔH_max (dépression admissible)", f"{fl['delta_h_max']:.2f} mCE"),
            ("P_nominale", f"{fl['p_nominale_bar']:.1f} bar = {fl['p_collapse_mce']:.1f} mCE"),
            ("Marge sécurité (50%)", f"{fl['marge_mce']:.1f} mCE"),
            ("Verdict", "CONFORME" if fl['conforme'] else "NON CONFORME"),
        ]:
            r = t_fl.add_row().cells; r[0].text = lib; r[1].text = val
        doc.add_paragraph(fl['message'])
        p = doc.add_paragraph(); p.add_run(f"Référence : {fl['reference']}").italic = True
    except Exception as e:
        doc.add_paragraph(f"(Vérification flambement indisponible : {e})")

    # 4.2 Perte de charge locale
    doc.add_heading("Perte de charge locale — Passage d'air à la vanne", level=2)
    doc.add_paragraph(
        "La perte de charge subie par l'air en traversant la vanne de sectionnement "
        "est estimée par la formule de perte locale : Δp = K × ρ_air × V² / 2")
    try:
        pc = etat.perte_charge_vanne()
        t_pc = doc.add_table(rows=1, cols=2)
        t_pc.style = "Light Grid Accent 1"
        t_pc.rows[0].cells[0].text = "Paramètre"; t_pc.rows[0].cells[1].text = "Valeur"
        for lib, val in [
            ("K (coefficient de perte)", f"{pc['k']:.1f} (vanne {'ouverte' if pc['k'] < 1 else 'fermée'})"),
            ("ρ_air", f"{pc['rho_air']:.3f} kg/m³"),
            ("V_collet", f"{pc['v_collet']:.1f} m/s"),
            ("Δp", f"{pc['delta_p_pa']:.1f} Pa = {pc['delta_p_mmceau']:.2f} mmCE eau"),
            ("Vitesse limite", f"{pc['vitesse_limite']:.0f} m/s (ISO 6358-1)"),
            ("Verdict", "CONFORME" if pc['conforme'] else "NON CONFORME"),
        ]:
            r = t_pc.add_row().cells; r[0].text = lib; r[1].text = val
        doc.add_paragraph(pc['message'])
        p = doc.add_paragraph(); p.add_run(f"Référence : {pc['reference']}").italic = True
    except Exception as e:
        doc.add_paragraph(f"(Calcul perte de charge indisponible : {e})")

    doc.add_paragraph()
    # --- ANNEXE : catalogue des équipements (vérification variantes conformes) ---
    doc.add_page_break()
    doc.add_heading("Annexe — Catalogue des équipements (vérification)", level=1)
    doc.add_paragraph(
        "Catalogue des organes d'air par fournisseur et rôle, avec leur capacité "
        "d'air admise (m³/h à la dépression nominale de 3 mCE, NF EN 805). Les "
        "diamètres EFFECTIVEMENT RETENUS dans la proposition / variante / groupe "
        "sont signalés « retenu (xN) ». Ce catalogue permet de vérifier "
        "l'adéquation des variantes conformes — TRIFON et CEAI sont présentés "
        "en deux entrées distinctes pour faciliter le contrôle.")
    try:
        for sec in etat.catalogue_verification():
            doc.add_heading(
                f"{sec['nom']} — {sec['role_lib']}", level=2)
            cpt = doc.add_paragraph().add_run(
                f"Statut : {sec['statut'] or '—'}  ·  Origine : {sec['origine'] or '—'}")
            cpt.italic = True
            t = doc.add_table(rows=1, cols=3)
            t.style = "Light Grid Accent 1"
            t.rows[0].cells[0].text = "DN"
            t.rows[0].cells[1].text = "Capacité (m³/h)"
            t.rows[0].cells[2].text = "Retenu"
            for it in sec["items"]:
                cells = t.add_row().cells
                cells[0].text = str(it["dn"])
                cells[1].text = f"{it['q_capacity']:,.0f}"
                cells[2].text = (f"OUI (×{it['nombre']})" if it["retenu"]
                                 else "—")
    except Exception as e:
        doc.add_paragraph(f"(Annexe catalogue indisponible : {e})")

    doc.add_paragraph()
    doc.add_paragraph("Fin du rapport — Document établi avec l'outil de dimensionnement "
                      "(méthode de calcul MK_A.A 2026 appliquée) "
                      f"({datetime.date.today().isoformat()} {_timestamp()})", style="Footer")

    doc.save(chemin)
    return chemin


# ---------------------------------------------------------------------------
# .XLSX
# ---------------------------------------------------------------------------
def exporter_xlsx(etat, chemin: str) -> str:
    """Génère un classeur .xlsx avec les feuilles : Synthèse, Calcul, Dimensionnement,
    Vérification, Catalogue."""
    etat.calculer()  # le rapport doit toujours refléter l'état courant
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    titre_font = Font(bold=True, size=13)
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F6FEB")
    ok_fill = PatternFill("solid", fgColor="2ECC71")
    nok_fill = PatternFill("solid", fgColor="E74C3C")

    def _entete(ws, entetes):
        for j, htxt in enumerate(entetes, start=1):
            c = ws.cell(row=1, column=j, value=htxt)
            c.font = hdr_font
            c.fill = hdr_fill
            c.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"

    def _autowidth(ws, ncols, maxw=45):
        for j in range(1, ncols + 1):
            m = 0
            for row in ws.iter_rows(min_col=j, max_col=j):
                for cell in row:
                    if cell.value is not None:
                        m = max(m, min(len(str(cell.value)), maxw))
            ws.column_dimensions[get_column_letter(j)].width = m + 2

    # --- Feuille Synthèse ---
    ws = wb.active
    ws.title = "Synthèse"
    ws["A1"] = _titre_rapport(etat); ws["A1"].font = titre_font
    for i, (lib, val) in enumerate(_info_projet(etat), start=3):
        ws.cell(row=i, column=1, value=lib).font = Font(bold=True)
        ws.cell(row=i, column=2, value=val)
    ws.cell(row=10, column=1, value="Conduite (DN)").font = Font(bold=True)
    dc = etat.dc()
    ws.cell(row=10, column=2, value=f"DN {dc.dn:g} mm, D = {dc.d:.3f} m, S = {dc.section:.4f} m²")
    ws.cell(row=11, column=1, value="Q_Ve (débit admis)").font = Font(bold=True)
    ws.cell(row=11, column=2, value=round(etat.resultat.q_ve_m3h, 1))
    r = 13
    ws.cell(row=r, column=1, value="RÉFÉRENCES NORMATIVES").font = Font(bold=True, size=11)
    for ref in REFERENCES_NORMATIVES:
        r += 1
        ws.cell(row=r, column=1, value="•")
        ws.cell(row=r, column=2, value=ref)
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 60

    # --- Feuille Calcul ---
    ws = wb.create_sheet("Calcul")
    _entete(ws, ["Étape", "Détail"])
    for i, step in enumerate(etat.resultat.journal, start=2):
        ws.cell(row=i, column=1, value=f"Étape {i-1}")
        ws.cell(row=i, column=2, value=step)
    _autowidth(ws, 2)

    # --- Feuille Dimensionnement ---
    # Les organes théoriques ne sont plus dimensionnés dans le rapport : la
    # proposition client est vérifiée directement (Fs sans majoration).
    ws = wb.create_sheet("Dimensionnement")
    _entete(ws, ["Donnée", "Valeur"])
    ws.cell(row=2, column=1, value="Note")
    ws.cell(row=2, column=2,
            value=("Dimensionnement théorique omis : la proposition client est "
                   "vérifiée DIRECTEMENT en besoins bruts sans majoration "
                   "(feuille « Vérification ») — Fs = Capacité / Demande ; le "
                   "contrôleur juge la valeur du Fs dans ses écrits."))
    ws.cell(row=2, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    row = 4
    fv = etat.derniere_trace.get("fournisseur_ventouse") or ""
    fc = etat.derniere_trace.get("fournisseur_clapet") or ""
    ws.cell(row=row, column=1, value="Fournisseur ventouse")
    ws.cell(row=row, column=3, value=fv or "TRIFON (standard)")
    ws.cell(row=row + 1, column=1, value="Fournisseur clapet")
    ws.cell(row=row + 1, column=3, value=fc or "CEAI (standard)")
    ws.cell(row=row + 2, column=1, value="Remplissage (§9)").font = Font(bold=True)
    remp = etat.derniere_trace["remplissage"]
    ws.cell(row=row + 3, column=1, value="Débit remplissage (m³/h)")
    ws.cell(row=row + 3, column=2, value=round(remp["q"], 1))
    p = remp["purgeurs"]
    ws.cell(row=row + 4, column=1, value="Purgeur PSA")
    ws.cell(row=row + 4, column=3, value=f"{p.nombre} × {p.nom} DN{p.dn}")
    sn = remp.get("purgeurs_snh")
    if sn is not None:
        ws.cell(row=row + 5, column=1, value="Purgeur sonic SNH (NSH)")
        ws.cell(row=row + 5, column=3, value=f"{sn.nombre} × {sn.nom} DN{sn.dn}")
        ws.cell(row=row + 6, column=1, value="Q_fill sonique NSH (m³/h, P_fill/P_svc)")
        ws.cell(row=row + 5, column=2, value=round(sn.capacite_unitaire_m3h, 1))
    _autowidth(ws, 6)

    # --- Feuille Vérification ---
    ws = wb.create_sheet("Vérification")
    _entete(ws, ["Repére", "Demande (m³/h)", "Catégorie", "DN (mm)", "Nbr",
                 "Implantation", "Fournisseur", "Capacité (m³/h)", "Fs", "Verdict"])
    for i, r in enumerate(etat.tableau_verification_organes(), start=2):
        ws.cell(row=i, column=1, value=r["repere"])
        ws.cell(row=i, column=2, value=(round(r["demande_m3h"], 2)
                                        if r["demande_m3h"] is not None else None))
        ws.cell(row=i, column=3, value=r["categorie"])
        ws.cell(row=i, column=4, value=(r["dn"] if r["dn"] is not None else "—"))
        ws.cell(row=i, column=5, value=r["nombre"])
        ws.cell(row=i, column=6, value=r["implantation"])
        ws.cell(row=i, column=7, value=r["fournisseur"])
        ws.cell(row=i, column=8, value=(round(r["capacite_m3h"], 2)
                                        if r["capacite_m3h"] is not None else None))
        ws.cell(row=i, column=9, value=(r["fs"] if r["fs"] is not None else None))
        vcell = ws.cell(row=i, column=10, value=r["verdict"])
        vcell.fill = ok_fill if r["verdict"] == "Conforme" else nok_fill
        vcell.font = Font(bold=True, color="FFFFFF")
    _autowidth(ws, 10)

    # --- Feuille Implantation (amont/aval) ---
    ws = wb.create_sheet("Implantation")
    ws["A1"] = "Implantation & justification amont/aval"; ws["A1"].font = Font(bold=True, size=12)
    plan = etat.plan_pose()
    ws["A3"] = plan["repere"]; ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A3:B3")
    ws["A4"] = "Position (par rapport à la vanne de sectionnement)"; ws["A4"].font = Font(bold=True)
    ws["B4"] = "Organes"; ws["B4"].font = Font(bold=True)
    rr = 5
    for pos, lib in (("amont", "AMONT — côté point haut Ve"),
                     ("aval", "AVAL — côté conduite"),
                     ("pi1", "PI1 — point intermédiaire amont (organes dédiés)"),
                     ("pi2", "PI2 — point intermédiaire aval (organes dédiés)"),
                     ("sans_position", "Sans position")):
        items = plan["organes_par_position"].get(pos, [])
        ws.cell(row=rr, column=1, value=lib)
        ws.cell(row=rr, column=2, value=" ; ".join(items) if items else "—")
        ws.cell(row=rr, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        rr += 1
    if plan.get("combo"):
        rr += 1
        ws.cell(row=rr, column=1, value="Justification").font = Font(bold=True)
        for j in plan["combo"].get("justifications", []):
            rr += 1
            ws.cell(row=rr, column=1, value="•")
            ws.cell(row=rr, column=2, value=j)
            ws.cell(row=rr, column=2).alignment = Alignment(wrap_text=True, vertical="top")
    _autowidth(ws, 2)

    # --- Feuille Vérifications (flambement + perte de charge) ---
    ws = wb.create_sheet("Vérifications")
    ws["A1"] = "Vérifications structurelles et aérauliques"; ws["A1"].font = Font(bold=True, size=12)
    rr = 3
    # 4.1 Flambement
    ws.cell(row=rr, column=1, value="4.1 Résistance au vide — Flambement").font = Font(bold=True)
    rr += 1
    ws.cell(row=rr, column=1, value="Critère"); ws.cell(row=rr, column=2, value="ΔH_max < 50% × P_nominale (mCE)")
    rr += 1
    try:
        fl = etat.verifier_flambement()
        for lib, val in [
            ("ΔH_max", f"{fl['delta_h_max']:.2f} mCE"),
            ("P_nominale", f"{fl['p_nominale_bar']:.1f} bar = {fl['p_collapse_mce']:.1f} mCE"),
            ("Marge sécurité (50%)", f"{fl['marge_mce']:.1f} mCE"),
            ("Verdict", "CONFORME" if fl['conforme'] else "NON CONFORME"),
            ("Référence", fl['reference']),
        ]:
            ws.cell(row=rr, column=1, value=lib); ws.cell(row=rr, column=2, value=val)
            rr += 1
    except Exception as e:
        ws.cell(row=rr, column=1, value=f"(Indisponible : {e})"); rr += 1
    rr += 1
    # 4.2 Perte de charge
    ws.cell(row=rr, column=1, value="4.2 Perte de charge locale — Vanne").font = Font(bold=True)
    rr += 1
    ws.cell(row=rr, column=1, value="Formule"); ws.cell(row=rr, column=2, value="Δp = K × ρ_air × V² / 2")
    rr += 1
    try:
        pc = etat.perte_charge_vanne()
        for lib, val in [
            ("K", f"{pc['k']:.1f}"),
            ("ρ_air", f"{pc['rho_air']:.3f} kg/m³"),
            ("V_collet", f"{pc['v_collet']:.1f} m/s"),
            ("Δp", f"{pc['delta_p_pa']:.1f} Pa = {pc['delta_p_mmceau']:.2f} mmCE eau"),
            ("Vitesse limite", f"{pc['vitesse_limite']:.0f} m/s (ISO 6358-1)"),
            ("Verdict", "CONFORME" if pc['conforme'] else "NON CONFORME"),
            ("Référence", pc['reference']),
        ]:
            ws.cell(row=rr, column=1, value=lib); ws.cell(row=rr, column=2, value=val)
            rr += 1
    except Exception as e:
        ws.cell(row=rr, column=1, value=f"(Indisponible : {e})"); rr += 1
    rr += 1
    # 4.3 Dimensionnement corrigé — groupe retenu
    ws.cell(row=rr, column=1, value="4.3 Dimensionnement corrigé — groupe retenu").font = Font(bold=True)
    rr += 1
    try:
        comp = etat.proposition_client_justifiee()
    except Exception:
        comp = None
    if comp is not None:
        rr += 1
        ws.cell(row=rr, column=1, value="Tableau 2 — Dimensionnement corrigé (groupe retenu)").font = Font(bold=True)
        rr += 1
        ws.cell(row=rr, column=1, value=comp["verdict_global"])
        rr += 2
        for j, htxt in enumerate(["Rép.", "Catégorie", "DN (mm)", "Nbr", "Fournisseur",
                                  "Rôle / capacité"], start=1):
            c = ws.cell(row=rr, column=j, value=htxt)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        rr += 1
        for o in comp["retenu"]:
            ws.cell(row=rr, column=1, value=o["rep"])
            ws.cell(row=rr, column=2, value=o["categorie"])
            ws.cell(row=rr, column=3, value=(o["dn"] if o["dn"] is not None else "—"))
            ws.cell(row=rr, column=4, value=o["nombre"])
            ws.cell(row=rr, column=5, value=o["fournisseur"] or "—")
            ws.cell(row=rr, column=6, value=o["role"])
            ws.cell(row=rr, column=6).alignment = Alignment(wrap_text=True, vertical="top")
            rr += 1
    _autowidth(ws, 7)

    # --- Feuille Plan X, Y, Z (exécution) ---
    ws = wb.create_sheet("Plan XYZ")
    ws["A1"] = "Plan d'implantation — coordonnées X, Y, Z (exécution)"; ws["A1"].font = Font(bold=True, size=12)
    xyz = etat.plan_xyz()
    ws["A3"] = xyz["base"]; ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A3:I3")
    entetes_xyz = ["Rép.", "Organe", "DN", "Nbr", "Position", "X (m)", "Y (m)", "Z (m)", "Remarque"]
    for j, h in enumerate(entetes_xyz, start=1):
        c = ws.cell(row=4, column=j, value=h); c.font = Font(bold=True)
    rrr = 5
    for o in xyz["organes"]:
        ws.cell(row=rrr, column=1, value=o["ref"])
        ws.cell(row=rrr, column=2, value=o["type"])
        ws.cell(row=rrr, column=3, value=f"DN{o['dn']}")
        ws.cell(row=rrr, column=4, value=o["nombre"])
        ws.cell(row=rrr, column=5, value=o["position"])
        ws.cell(row=rrr, column=6, value=o["x_m"])
        ws.cell(row=rrr, column=7, value=o["y_m"])
        ws.cell(row=rrr, column=8, value=o["z_m"])
        ws.cell(row=rrr, column=9, value=o["remarque"])
        ws.cell(row=rrr, column=9).alignment = Alignment(wrap_text=True, vertical="top")
        rrr += 1
    _autowidth(ws, 8)

    # --- Feuille Catalogue ---
    from ..data.catalogues import TRIFON, CEAI, PSA, DEPRESSION_NOMINALE
    from ..data import valve_db
    ws = wb.create_sheet("Catalogue")
    ws["A1"] = "Catalogue des équipements — vérification des variantes"
    ws["A1"].font = titre_font
    try:
        cats = etat.catalogue_verification()
        if cats:
            _entete(ws, ["Fournisseur", "Rôle", "DN", "Capacité (m³/h)",
                         "Retenu", "Statut", "Origine"])
            row = 2
            for sec in cats:
                four_font = Font(bold=True)
                for it in sec["items"]:
                    ws.cell(row=row, column=1, value=sec["nom"]).font = four_font
                    ws.cell(row=row, column=2, value=sec["role_lib"])
                    ws.cell(row=row, column=3, value=it["dn"])
                    ws.cell(row=row, column=4, value=round(it["q_capacity"], 1))
                    rc = ws.cell(row=row, column=5,
                                 value=(f"OUI (×{it['nombre']})" if it["retenu"] else "—"))
                    if it["retenu"]:
                        rc.fill = ok_fill
                    ws.cell(row=row, column=6, value=sec["statut"] or "—")
                    ws.cell(row=row, column=7, value=sec["origine"] or "—")
                    row += 1
        else:
            raise RuntimeError("catalogue vide")
        _autowidth(ws, 7)
    except Exception:
        # Repli sur le catalogue interne (comportement historique)
        _entete(ws, ["Type", "DN", "Capacité @ -3 mce (m³/h)", "Remarque"])
        row = 2
        for dn, v in TRIFON.items():
            ws.cell(row=row, column=1, value="Ventouse TRIFON")
            ws.cell(row=row, column=2, value=dn)
            ws.cell(row=row, column=3, value=v["q"][DEPRESSION_NOMINALE])
            row += 1
        for dn, v in CEAI.items():
            ws.cell(row=row, column=1, value="Clapet CEAI")
            ws.cell(row=row, column=2, value=dn)
            ws.cell(row=row, column=3, value=v["q"][DEPRESSION_NOMINALE])
            row += 1
        for dn, v in PSA.items():
            ws.cell(row=row, column=1, value="Purgeur PSA")
            ws.cell(row=row, column=2, value=dn)
            ws.cell(row=row, column=3, value=v["q_remplissage"])
            ws.cell(row=row, column=4, value="Remplissage")
            row += 1
        _autowidth(ws, 4)
    ws.cell(row=row + 1, column=1,
            value=f"Source : {valve_db.source_base()}").font = Font(italic=True, color="808080")

    # --- Feuille Cumul des débits : casse franche (Brèche) OU Q MK_A.A ---
    if _utilise_breche(etat):
        _ecrire_cumul_breche_xlsx(etat, ws, wb, titre_font, ok_fill, nok_fill, _autowidth)
    else:
        _ecrire_debits_cumules_xlsx(etat, ws, wb, titre_font, ok_fill, nok_fill, _autowidth)

    wb.save(chemin)
    return chemin


# ---------------------------------------------------------------------------
# Note MÉTHODOLOGIE MK_A.A (2026) — document Word séparé du rapport
# ---------------------------------------------------------------------------
def exporter_methodologie(etat, chemin: str) -> str:
    """Génère la note méthodologique « MÉTHODOLOGIE MK_A.A_2026 » (.docx).

    Document autonome, séparé du rapport de synthèse, décrivant les phases de
    calcul de la méthode explicite MK_A.A (2026), destiné au contrôleur.
    """
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # En-tête : icône + référence (identique au rapport)
    try:
        hp = doc.sections[0].header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        icone_png = _chemin_png_icone()
        if icone_png:
            run_img = hp.add_run()
            try:
                from PIL import Image as _PILImage
                tmp_ico = os.path.join(tempfile.gettempdir(),
                                       "_mk_aa_docx_icone_petite.png")
                with _PILImage.open(icone_png) as im:
                    im.convert("RGBA").resize((64, 64), _PILImage.LANCZOS).save(tmp_ico)
                run_img.add_picture(tmp_ico, height=Inches(0.38))
            except Exception:
                run_img.add_picture(icone_png, height=Inches(0.38))
            run_txt = hp.add_run("   MK_A.A 2026 — Débits d'air admis en conduites AEP")
        else:
            run_txt = hp.add_run("MK_A.A 2026 — Débits d'air admis en conduites AEP")
        run_txt.font.size = Pt(9)
        run_txt.font.color.rgb = RGBColor(0x1F, 0x6F, 0xEB)
    except Exception:
        pass

    h = doc.add_heading("MÉTHODOLOGIE MK_A.A_2026", level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp = doc.add_paragraph(
        "Phases de calcul de la méthode explicite MK_A.A (2026)")
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Identification", level=1)
    for lib, val in _info_projet(etat):
        p = doc.add_paragraph()
        p.add_run(f"{lib} : ").bold = True
        p.add_run(str(val))

    doc.add_heading("Objet", level=1)
    doc.add_paragraph(_METHODOLOGIE_INTRO)

    doc.add_heading("Références normatives", level=1)
    doc.add_paragraph("La méthode s'appuie sur les textes suivants :")
    for ref in REFERENCES_NORMATIVES:
        doc.add_paragraph(ref, style="List Bullet")

    doc.add_heading("Phases de calcul", level=1)
    for titre, corps in _PHASES_METHODOLOGIE:
        doc.add_heading(titre, level=2)
        doc.add_paragraph(corps)

    doc.add_paragraph()
    note_ctrl = doc.add_paragraph()
    note_ctrl.add_run("Note pour le contrôle : ").bold = True
    note_ctrl.add_run(_METHODOLOGIE_NOTE)

    doc.add_paragraph()
    doc.add_paragraph(
        f"Document établi avec l'outil de dimensionnement "
        f"(méthode de calcul MK_A.A 2026 appliquée) "
        f"({datetime.date.today().isoformat()} {_timestamp()})")
    doc.save(chemin)
    return chemin


# =====================================================================
# EXPORT NOMENCLATURE (XLSX)
# =====================================================================
def exporter_nomenclature(etat, chemin: str) -> str:
    """Exporte la nomenclature des organes et pièces annexes au format XLSX.

    Génère deux feuilles :
      - « Organes d'air » : trifon, clapets, purgeurs du groupe validé
      - « Pièces annexes » : vannes d'isolement, joints, vanne principale
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    def _aw(ws, ncols, maxw=45):
        for j in range(1, ncols + 1):
            m = 0
            for row in ws.iter_rows(min_col=j, max_col=j):
                for cell in row:
                    v = str(cell.value or "")
                    m = max(m, min(len(v), maxw))
            ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = m + 3

    wb = Workbook()
    ws = wb.active
    ws.title = "Organes d'air"
    ws["A1"] = "Nomenclature des organes — Groupe d'air validé / conforme"
    ws["A1"].font = Font(bold=True, size=12)

    # Police, couleurs
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill(start_color="4472C4", end_color="4472C4",
                           fill_type="solid")

    nom = etat.nomenclature()
    ws["A3"] = nom["base"]
    ws["A3"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A3:F3")

    # --- Feuille 1 : Organes d'air ---
    heads = ["Rép.", "Désignation", "DN (mm)", "PN (bar)", "Qté",
             "Position & Rôle Technique"]
    for j, h in enumerate(heads, start=1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = hdr_font; c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A5"

    rr = 5
    for o in nom["organes"]:
        ws.cell(row=rr, column=1, value=o["rep"])
        ws.cell(row=rr, column=2, value=o["designation"])
        ws.cell(row=rr, column=3, value=o["dn"])
        ws.cell(row=rr, column=4, value=o["pn"])
        ws.cell(row=rr, column=5, value=o["qte"])
        ws.cell(row=rr, column=6, value=o["role"])
        ws.cell(row=rr, column=6).alignment = Alignment(wrap_text=True,
                                                         vertical="top")
        rr += 1

    _aw(ws, 6)

    # --- Feuille 2 : Pièces annexes ---
    ws2 = wb.create_sheet("Pièces annexes")
    ws2["A1"] = "Pièces annexes — Vannes d'isolement, joints, vanne principale"
    ws2["A1"].font = Font(bold=True, size=12)
    heads2 = ["Rép.", "Désignation", "DN (mm)", "PN (bar)", "Qté", "Rôle"]
    for j, h in enumerate(heads2, start=1):
        c = ws2.cell(row=3, column=j, value=h)
        c.font = hdr_font; c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center")
    ws2.freeze_panes = "A4"

    rr = 4
    for i, c in enumerate(nom["complement"], start=1):
        ws2.cell(row=rr, column=1, value=f"{i:02d}")
        ws2.cell(row=rr, column=2, value=c["designation"])
        ws2.cell(row=rr, column=3, value=c["dn"])
        ws2.cell(row=rr, column=4, value=c["pn"])
        ws2.cell(row=rr, column=5, value=c["qte"])
        ws2.cell(row=rr, column=6, value=c["role"])
        ws2.cell(row=rr, column=6).alignment = Alignment(wrap_text=True,
                                                         vertical="top")
        rr += 1

    _aw(ws2, 6)
    wb.save(chemin)
    return chemin


# ---------------------------------------------------------------------------
# Rapport « profil complet » (chaîne de tronçons) — .docx et .xlsx
# ---------------------------------------------------------------------------

def exporter_docx_profil(pc, etat, chemin: str) -> str:
    """Génère le rapport .docx de la méthode « profil complet » (multi-TR)."""
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # En-tête (comme exporter_docx)
    try:
        hp = doc.sections[0].header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        icone_png = _chemin_png_icone()
        if icone_png:
            run_img = hp.add_run()
            try:
                from PIL import Image as _PILImage
                tmp_ico = os.path.join(tempfile.gettempdir(),
                                       "_mk_aa_docx_icone_petite.png")
                with _PILImage.open(icone_png) as im:
                    im.convert("RGBA").resize((64, 64), _PILImage.LANCZOS).save(tmp_ico)
                run_img.add_picture(tmp_ico, height=Inches(0.38))
            except Exception:
                run_img.add_picture(icone_png, height=Inches(0.38))
            run_txt = hp.add_run("   MK_A.A 2026 — Profil complet")
        else:
            run_txt = hp.add_run("MK_A.A 2026 — Profil complet")
        run_txt.font.size = Pt(9)
        run_txt.font.color.rgb = RGBColor(0x1F, 0x6F, 0xEB)
    except Exception:
        pass

    h = doc.add_heading(_titre_rapport(etat), level=0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        "Note technique MK_A.A 2026 — Méthode « profil complet » "
        "(chaîne de tronçons)").alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("Identification du projet", level=1)
    for lib, val in _info_projet(etat):
        p = doc.add_paragraph()
        p.add_run(f"{lib} : ").bold = True
        p.add_run(str(val))

    doc.add_heading("Présentation de la chaîne", level=1)
    p = doc.add_paragraph()
    p.add_run("Tronçons : ").bold = True
    p.add_run(" → ".join(rt.label for rt in pc.troncons) if pc.troncons else "—")
    p = doc.add_paragraph()
    p.add_run("Conduite : ").bold = True
    p.add_run(f"DN {pc.dn_mm:g} mm · T = {pc.temperature:.1f} °C · "
              f"PN {pc.pression_nominale:g} bar · "
              f"longueur développée totale = {pc.longueur_totale:,.0f} m")
    p = doc.add_paragraph()
    p.add_run("Méthode : ").bold = True
    p.add_run("chaque tronçon est calculé par la formule explicite MK_A.A "
              "(2026) ; les points Vi2(TRi) ≡ Vi1(TR(i+1)) sont recousus en "
              "vidange commune (débit d'évacuation = max des pentes adjacentes) "
              "; la continuité du profil est vérifiée.")

    # --- §1 Entrées par tronçon ---
    doc.add_heading("1. Entrées par tronçon", level=1)
    t = doc.add_table(rows=1, cols=9)
    t.style = "Light Grid Accent 1"
    for j, h in enumerate(["TR", "Schéma", "Z Vi1", "Z Ve", "Z Vi2", "a", "b",
                           "c", "d"]):
        t.rows[0].cells[j].text = h
    for rt in pc.troncons:
        tc = rt.troncon
        r = t.add_row().cells
        for j, v in enumerate([tc.label, tc.schema, tc.z_vi1, tc.z_ve, tc.z_vi2,
                               tc.a, tc.b, tc.c, tc.d]):
            r[j].text = f"{v:g}" if isinstance(v, float) else str(v)

    # --- §2 Résultats par tronçon ---
    doc.add_heading("2. Résultats par tronçon", level=1)
    t = doc.add_table(rows=1, cols=6)
    t.style = "Light Grid Accent 1"
    for j, h in enumerate(["TR", "Q_Ve", "Q amont", "Q aval", "Q_PI1", "Q_PI2"]):
        t.rows[0].cells[j].text = h
    q_max = 0.0
    for rt in pc.troncons:
        r = t.add_row().cells
        rs = rt.res
        q_max = max(q_max, rs.q_ve_m3h)
        for j, v in enumerate([rt.label, rs.q_ve_m3h, rs.q_am_m3h, rs.q_av_m3h,
                               rs.q_pi1_m3h, rs.q_pi2_m3h]):
            r[j].text = f"{v:,.0f}" if isinstance(v, (int, float)) and j else str(v)
    p = doc.add_paragraph()
    p.add_run(f"Débit maximal le long du profil : Q_Ve,max = {q_max:,.1f} m³/h.")

    # --- §3 Dimensionnement par tronçon ---
    doc.add_heading("3. Dimensionnement des organes (marge +15 %, plafond 90 %)",
                    level=1)
    for rt in pc.troncons:
        doc.add_heading(f"Tronçon {rt.label} — Q_Ve = {rt.res.q_ve_m3h:,.0f} m³/h "
                        f"(besoin majoré {rt.organes['ventouse'].besoin_m3h:,.0f} m³/h)",
                        level=2)
        t = doc.add_table(rows=1, cols=6)
        t.style = "Light Grid Accent 1"
        for j, h in enumerate(["Type", "Référence", "DN", "Qté", "Cap. inst.",
                               "Taux util."]):
            t.rows[0].cells[j].text = h
        for nom, sel in rt.organes.items():
            r = t.add_row().cells
            r[0].text = "Ventouse" if nom == "ventouse" else "Clapet d'admission"
            r[1].text = sel.nom
            r[2].text = f"DN{sel.dn}"
            r[3].text = str(sel.nombre)
            r[4].text = f"{sel.capacite_installee_m3h:,.0f} m³/h"
            r[5].text = f"{sel.taux_utilisation:.1f} %"
        if rt.remplissage.get("q"):
            p = doc.add_paragraph()
            p.add_run("Remplissage (§9) : ").bold = True
            pr = rt.remplissage["purgeurs"]
            p.add_run(f"Q = {rt.remplissage['q']:,.0f} m³/h → "
                      f"{pr.nombre} × {pr.nom} DN{pr.dn}.")

    # --- §4 Croquis du profil complet (PNG) ---
    doc.add_heading("4. Croquis du profil en long complet", level=1)
    try:
        chemin_png = os.path.join(tempfile.gettempdir(),
                                  "_mk_aa_profil_complet.png")
        croquis_mod.croquis_chain_png(pc, chemin_png)
        doc.add_picture(chemin_png, width=Inches(6.9))
        doc.add_paragraph("Points V (ventouse) · s (PI) · o (point bas). "
                          "Les jonctions sont cerclées en bleu.")
    except Exception as e:
        doc.add_paragraph(f"(Croquis PNG indisponible : {e})")
        ascii_c = croquis_mod.croquis_chain_ascii(pc)
        for ligne in ascii_c.splitlines():
            p = doc.add_paragraph()
            p.add_run(ligne).font.name = "Consolas"
            p.paragraph_format.space_after = Pt(0)

    # --- §5 Synthèse organes + vidanges ---
    doc.add_heading("5. Synthèse des organes le long du profil", level=1)
    t = doc.add_table(rows=1, cols=9)
    t.style = "Light Grid Accent 1"
    for j, h in enumerate(["Rép.", "TR", "Point", "Type", "DN", "Qté",
                           "Besoin (m³/h)", "Z (m)", "x (m)"]):
        t.rows[0].cells[j].text = h
    for o in pc.synthese_organes():
        r = t.add_row().cells
        r[0].text = o["rep"]
        r[1].text = o["tr"]
        r[2].text = o["point"]
        r[3].text = o["type"]
        r[4].text = f"DN{o['dn']}"
        r[5].text = str(o["nombre"])
        r[6].text = f"{o['besoin_m3h']:,.0f}"
        r[7].text = (f"{o['z']:.1f}" if o.get("z") is not None else "—")
        r[8].text = (f"{o['x']:,.0f}" if o.get("x") is not None else "—")

    vid = pc.synthese_vidanges()
    if vid:
        doc.add_heading("5bis. Vidanges communes (points bas)", level=1)
        t = doc.add_table(rows=1, cols=4)
        t.style = "Light Grid Accent 1"
        for j, h in enumerate(["Vidange", "Z (m NGM)", "Q évacuation",
                               "Continu (ΔZ ≈ 0)"]):
            t.rows[0].cells[j].text = h
        for v in vid:
            r = t.add_row().cells
            r[0].text = v["point"]
            r[1].text = f"{v['z']:.1f}"
            r[2].text = f"{v['q_evacuation']:,.0f} m³/h"
            r[3].text = "OUI" if v["continue"] else \
                "NON — vérifier la continuité du profil"

    # --- §6 Vérification proposition client ---
    if getattr(pc, "propositions", None) and any(pc.propositions):
        doc.add_heading("6. Vérification de la proposition client", level=1)
        doc.add_paragraph("Besoin majoré = Q_Ve × 1,15 (marge +15 %, plafond "
                          "90 %) — défauts v3. Verdict par organe et groupe "
                          "d'air de chaque tronçon.")
        t = doc.add_table(rows=1, cols=6)
        t.style = "Light Grid Accent 1"
        for j, h in enumerate(["TR", "Catégorie", "Proposé par le client",
                               "Besoin (m³/h)", "Capacité proposée (m³/h)",
                               "Verdict"]):
            t.rows[0].cells[j].text = h
        if pc.verifications:
            for v in pc.verifications:
                r = t.add_row().cells
                r[0].text = v["troncon"]
                r[1].text = v["categorie"]
                r[2].text = v["propose"]
                r[3].text = f"{v['besoin_m3h']:,.0f}"
                r[4].text = f"{v['cap_proposee_m3h']:,.0f}"
                r[5].text = "CONFORME" if v["conforme"] else "NON CONFORME"
        concl = pc.conclusion_globale()
        p = doc.add_paragraph()
        p.add_run("Conclusion globale : ").bold = True
        if concl["conforme"]:
            p.add_run(f"PROPOSITION CONFORME — {concl['nb_sains']}/"
                      f"{concl['nb_verdicts']} vérifications OK.")
        else:
            p.add_run(f"PROPOSITION NON CONFORME — {concl['nb_sains']}/"
                      f"{concl['nb_verdicts']} OK, {concl['nb_defauts']} à revoir.")
            for tr_l, vs in concl["defauts_par_troncon"].items():
                for v in vs:
                    pp = doc.add_paragraph(style="List Bullet")
                    pp.add_run(f"{tr_l} — {v['categorie']} : déficit "
                               f"{v['deficit_m3h']:,.0f} m³/h (proposé : "
                               f"{v['propose']}). ").italic = True
                    pp.add_run(v["message"])

        # --- §6 bis Récapitulatif implantation par point haut (Fs) ---
        doc.add_heading("6 bis. Récapitulatif implantation par point haut",
                        level=1)
        doc.add_paragraph("Fs = capacité installée / besoin Q (brut). "
                          "Fs ≥ 1,20 → CONFORME ; 1,00 ≤ Fs < 1,20 → "
                          "ADMISSIBLE — NON RECOMMANDÉ (+ suggestion) ; "
                          "Fs < 1,00 → NON CONFORME. Les purgeurs (dégazage) "
                          "figurent dans l'implantation mais sont exclus du "
                          "Fs d'admission.")
        t2 = doc.add_table(rows=1, cols=8)
        t2.style = "Light Grid Accent 1"
        for j, h in enumerate(["TR", "Point haut", "Implantation",
                               "Capacité (m³/h)", "Besoin Q (m³/h)", "Fs",
                               "Avis", "Suggestion"]):
            t2.rows[0].cells[j].text = h
        for r in pc.recap_ph():
            c2 = t2.add_row().cells
            c2[0].text = r["tr"]
            c2[1].text = r["ph"]
            c2[2].text = r["implantation"]
            c2[3].text = f"{r['cap_m3h']:,.0f}"
            c2[4].text = f"{r['besoin_m3h']:,.0f}"
            c2[5].text = f"{r['fs']:.2f}"
            c2[6].text = r["avis"]
            c2[7].text = r["suggestion"]

    # --- §6 bis Avertissements / cohérence ---
    if pc.avertissements:
        doc.add_heading("7. Cohérence de la chaîne — avertissements", level=1)
        for a in pc.avertissements:
            doc.add_paragraph(a, style="List Bullet")

    doc.add_heading("Normes et références techniques", level=1)
    for ref in REFERENCES_NORMATIVES:
        doc.add_paragraph(ref, style="List Bullet")

    doc.save(chemin)
    return chemin


def exporter_xlsx_profil(pc, etat, chemin: str) -> str:
    """Génère le classeur .xlsx « profil complet » (multi-TR).

    Feuilles : Synthèse organes · Résultats TR · Vidanges · Entrées.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    titre_font = Font(bold=True, size=13)
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F6FEB")

    def _autowidth(ws, ncols, maxw=50):
        from openpyxl.utils import get_column_letter
        for j in range(1, ncols + 1):
            m = 0
            for row in ws.iter_rows(min_col=j, max_col=j):
                for cell in row:
                    if cell.value is not None:
                        m = max(m, min(len(str(cell.value)), maxw))
            ws.column_dimensions[get_column_letter(j)].width = m + 2

    def _entete(ws, entetes, row=1):
        for j, htxt in enumerate(entetes, start=1):
            c = ws.cell(row=row, column=j, value=htxt)
            c.font = hdr_font
            c.fill = hdr_fill
            c.alignment = Alignment(horizontal="center")
        ws.freeze_panes = ws.cell(row=row + 1, column=1).coordinate

    # --- Synthèse organes ---
    ws = wb.active
    ws.title = "Synthèse organes"
    ws["A1"] = "Synthèse des organes — Profil complet" 
    ws["A1"].font = titre_font
    r = 3
    ws.cell(row=r, column=1, value="Tronçons").font = Font(bold=True)
    ws.cell(row=r, column=2,
            value=" → ".join(rt.label for rt in pc.troncons) if pc.troncons else "—")
    r += 1
    ws.cell(row=r, column=1, value="Conduite").font = Font(bold=True)
    ws.cell(row=r, column=2,
            value=f"DN {pc.dn_mm:g} mm · T = {pc.temperature:.1f} °C · "
                  f"PN {pc.pression_nominale:g} bar · Ltot = "
                  f"{pc.longueur_totale:,.0f} m")
    r += 1
    ws.cell(row=r, column=1, value="Méthode").font = Font(bold=True)
    ws.cell(row=r, column=2,
            value=f"Formule explicite MK_A.A 2026 · marge +15 % · "
                  f"plafond 90 % · vi2(TRi)≡vi1(TRi+1)")
    r += 2
    heads = ["Rép.", "TR", "Point", "Type", "DN", "Qté", "Besoin (m³/h)",
             "Z (m NGM)", "x cumulée (m)"]
    _entete(ws, heads, row=r)
    r += 1
    for o in pc.synthese_organes():
        ws.cell(row=r, column=1, value=o["rep"])
        ws.cell(row=r, column=2, value=o["tr"])
        ws.cell(row=r, column=3, value=o["point"])
        ws.cell(row=r, column=4, value=o["type"])
        ws.cell(row=r, column=5, value=o["dn"])
        ws.cell(row=r, column=6, value=o["nombre"])
        ws.cell(row=r, column=7, value=round(o["besoin_m3h"], 1))
        ws.cell(row=r, column=8, value=round(o["z"], 1) if o.get("z") else "—")
        ws.cell(row=r, column=9, value=round(o["x"]) if o.get("x") else "—")
        r += 1
    _autowidth(ws, len(heads))

    # --- Résultats TR ---
    ws2 = wb.create_sheet("Résultats TR")
    ws2["A1"] = "Résultats par tronçon (Q en m³/h)" 
    ws2["A1"].font = titre_font
    heads2 = ["TR", "Schéma", "Q_Ve", "Q amont", "Q aval", "Q_PI1", "Q_PI2"]
    _entete(ws2, heads2, row=3)
    r = 4
    for rt in pc.troncons:
        rs = rt.res
        vals = [rt.label, rs.cas, round(rs.q_ve_m3h, 1), round(rs.q_am_m3h, 1),
                round(rs.q_av_m3h, 1), round(rs.q_pi1_m3h, 1),
                round(rs.q_pi2_m3h, 1)]
        for j, v in enumerate(vals, start=1):
            ws2.cell(row=r, column=j, value=v)
        r += 1
    _autowidth(ws2, len(heads2))

    # --- Vidanges ---
    ws3 = wb.create_sheet("Vidanges")
    ws3["A1"] = "Vidanges du profil — points bas / jonctions communes"
    ws3["A1"].font = titre_font
    heads3 = ["Vidange", "Z (m NGM)", "Q amont TRi (m³/h)", "Q aval TRi+1 (m³/h)",
              "Q évacuation (m³/h)", "Continuité ΔZ"]
    _entete(ws3, heads3, row=3)
    r = 4
    for v in pc.synthese_vidanges():
        for j, val in enumerate([v["point"], round(v["z"], 1),
                                 round(v["q_amont_tr_i"], 1),
                                 round(v["q_aval_tr_i1"], 1),
                                 round(v["q_evacuation"], 1),
                                 "OK (ΔZ < 0,01 m)" if v["continue"]
                                 else f"ÉCART {v['ecart_m']:.2f} m"], start=1):
            ws3.cell(row=r, column=j, value=val)
        r += 1
    _autowidth(ws3, len(heads3))

    # --- Entrées tronçons ---
    ws4 = wb.create_sheet("Entrées tronçons")
    ws4["A1"] = "Entrées par tronçon (schéma, cotes, distances)"
    ws4["A1"].font = titre_font
    heads4 = ["TR", "Schéma", "Z Vi1", "Z PI1", "Z Ve", "Z PI2", "Z Vi2",
              "a", "b", "c", "d", "l1", "l2"]
    _entete(ws4, heads4, row=3)
    r = 4
    for rt in pc.troncons:
        t = rt.troncon
        vals = [t.label, t.schema, t.z_vi1, t.z_pi1, t.z_ve, t.z_pi2, t.z_vi2,
                t.a, t.b, t.c, t.d, t.l1, t.l2]
        for j, v in enumerate(vals, start=1):
            ws4.cell(row=r, column=j, value=v)
        r += 1
    _autowidth(ws4, len(heads4))

    # --- Vérification proposition client ---
    if getattr(pc, "propositions", None) and any(pc.propositions):
        ws5 = wb.create_sheet("Vérification client")
        ws5["A1"] = "Vérification proposition client vs besoins (+15 %, plafond 90 %)"
        ws5["A1"].font = titre_font
        heads5 = ["TR", "Catégorie", "Proposé", "Besoin (m³/h)",
                  "Capacité proposée (m³/h)", "Déficit (m³/h)", "Verdict"]
        _entete(ws5, heads5, row=3)
        r = 4
        for v in pc.verifications:
            vals5 = [v["troncon"], v["categorie"], v["propose"],
                     round(v["besoin_m3h"], 1), round(v["cap_proposee_m3h"], 1),
                     round(v["deficit_m3h"], 1),
                     "CONFORME" if v["conforme"] else "NON CONFORME"]
            for j, val in enumerate(vals5, start=1):
                ws5.cell(row=r, column=j, value=val)
            r += 1
        concl = pc.conclusion_globale()
        ws5.cell(row=r + 1, column=1,
                 value="CONCLUSION : " + ("PROPOSITION CONFORME" if concl["conforme"]
                                          else "PROPOSITION NON CONFORME")).font = \
            Font(bold=True)
        ws5.cell(row=r + 2, column=1,
                 value=f"{concl['nb_sains']}/{concl['nb_verdicts']} vérifications "
                       f"OK ; {concl['nb_defauts']} à revoir.")

        # --- Récapitulatif implantation par point haut (Fs) ---
        rr = r + 4
        ws5.cell(row=rr, column=1,
                 value="Récapitulatif implantation par point haut "
                       "(Fs = capacité / besoin Q)").font = Font(bold=True)
        ws5.cell(row=rr + 1, column=1,
                 value="Fs ≥ 1,20 → CONFORME · 1,00 ≤ Fs < 1,20 → ADMISSIBLE — "
                       "NON RECOMMANDÉ (+ suggestion) · Fs < 1,00 → NON "
                       "CONFORME. Purgeurs (dégazage) exclus du Fs.")
        rr += 2
        for j, h in enumerate(["TR", "Point haut", "Implantation",
                               "Capacité (m³/h)", "Besoin Q (m³/h)", "Fs",
                               "Avis", "Suggestion"], start=1):
            cc = ws5.cell(row=rr, column=j, value=h)
            cc.font = Font(bold=True)
        rr += 1
        for row_ph in pc.recap_ph():
            vals = [row_ph["tr"], row_ph["ph"], row_ph["implantation"],
                    round(row_ph["cap_m3h"], 1), round(row_ph["besoin_m3h"], 1),
                    round(row_ph["fs"], 2), row_ph["avis"], row_ph["suggestion"]]
            for j, val in enumerate(vals, start=1):
                ws5.cell(row=rr, column=j, value=val)
            rr += 1
        _autowidth(ws5, len(heads5))

    wb.save(chemin)
    return chemin
