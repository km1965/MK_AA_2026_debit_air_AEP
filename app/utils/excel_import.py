"""Import d'un fichier Excel de données projet (modèle V05 — ROADMAP §14.1).

Le fichier modèle comporte deux feuilles :
  * « Données projet » : une ligne par paramètre
        A = Paramètre | B = Valeur | C = Unité | D = Description
    Les lignes d'en-tête de section commencent par « — » et sont ignorées.
  * « Organes client » : une ligne par organe
        A = Type | B = DN (mm) | C = Nombre | D = Position | E = Fournisseur

La fonction principale `lire_donnees_projet()` retourne :
  * un dict au format de `EtatApplication.to_dict()` — seuls les champs
    renseignés dans le fichier sont présents (les autres gardent leur valeur
    par défaut dans l'application) ;
  * une liste d'avertissements (contrôle de validité : valeurs manquantes,
    schéma inconnu, distances absentes pour le schéma, lignes d'organes
    incomplètes…).
"""

from __future__ import annotations

import unicodedata

import openpyxl

# ---------------------------------------------------------------------------
# Schémas de vidange acceptés (identiques à app.controller.SCHEMAS) et
# distances attendues par schéma (copie locale de controller.DISTANCES_PAR_CAS,
# évite un import circulaire controlleur ↔ utils).
# ---------------------------------------------------------------------------
ALLOWED_SCHEMAS = ("1A", "1B", "2A", "2B", "3A", "3B", "4A", "4B")

DISTANCES_PAR_CAS = {
    "1A": ("l1",),
    "1B": ("a", "b"),
    "2A": ("l2",),
    "2B": ("c", "d"),
    "3A": ("l1", "l2"),
    "3B": ("a", "b", "c", "d"),
    "4A": ("a", "b", "l2"),
    "4B": ("l1", "c", "d"),
}

_TYPES_AIR = ("trifon", "ceai", "psa")
# Alias acceptés dans la colonne « Type » de la feuille Organes client.
_TYPES_ALIASES = {
    "trifon": "trifon", "ventouse": "trifon", "ventouses": "trifon",
    "ceai": "ceai", "clapet": "ceai", "clapets": "ceai",
    "clapetadmission": "ceai", "psa": "psa", "purgeur": "psa",
    "vannesectionnement": "vanne", "vanne": "vanne",
}

# ---------------------------------------------------------------------------
# Correspondance libellés « Paramètre » → champs de l'état projet.
# Certaines valeurs sont déduites par préfixe (normalisé, trié du plus long
# au plus court) pour tolérer les variantes (« Diamètre nominal DN »,
# « Diamètre nominal DN (mm) », « Z_Ve (m) »…). Les libellés courts (« a »,
# « b », « L1 »…) sont validés en correspondance exacte uniquement.
# ---------------------------------------------------------------------------
_EXACT = {
    "a": "a",
    "b": "b",
    "c": "c",
    "d": "d",
    "l1": "l1",
    "l2": "l2",
    "schema": "schema",
    "schemadevidange": "schema",
}

_PREFIXES = (
    ("maitredouvrage", "moe"),
    ("projetmarche", "projet"),
    ("projet", "projet"),
    ("referencedocument", "reference"),
    ("branchesetudiees", "branches"),
    ("branches", "branches"),
    ("diametrenominaldn", "dn_mm"),
    ("temperaturedeleau", "temperature"),
    ("temperature", "temperature"),
    ("pressionnominalepn", "pression_nominale"),
    ("pressionnominale", "pression_nominale"),
    ("zvi1", "z_vi1"),
    ("zve", "z_ve"),
    ("zpi1", "z_pi1"),
    ("zpi2", "z_pi2"),
    ("zvi2", "z_vi2"),
    ("avi1pi1", "a"),
    ("bpi1ve", "b"),
    ("cvepi2", "c"),
    ("dpi2vi2", "d"),
    ("l1vi1ve", "l1"),
    ("l2vevi2", "l2"),
    ("hzdenivele", "hz_casse_franche"),
    ("dnvannesectionnement", "dn_vanne_sectionnement"),
    ("dnvanne", "dn_vanne_sectionnement"),
    ("vannesectionnement", "dn_vanne_sectionnement"),
)
_PREFIXES = tuple(sorted(_PREFIXES, key=lambda t: len(t[0]), reverse=True))

# Feuille « Données projet » : champs numériques.
_CHAMPS_NUM = {
    "dn_mm", "temperature", "pression_nominale",
    "z_vi1", "z_pi1", "z_ve", "z_pi2", "z_vi2",
    "a", "b", "c", "d", "l1", "l2",
    "hz_casse_franche", "dn_vanne_sectionnement",
}


# ---------------------------------------------------------------------------
def normaliser(s) -> str:
    """Minuscules sans accents, sans espaces ni ponctuation (clé de rapprochement)."""
    if s is None:
        return ""
    nfkd = unicodedata.normalize("NFKD", str(s))
    base = "".join(c for c in nfkd if not unicodedata.combining(c))
    return "".join(c for c in base.lower() if c.isalnum())


def _champ(libelle: str):
    """Renvoie la clé EtatApplication correspondant au libellé, sinon None."""
    n = normaliser(libelle)
    if not n:
        return None
    if n in _EXACT:
        return _EXACT[n]
    for prefix, champ in _PREFIXES:
        if n.startswith(prefix):
            return champ
    return None


def _parse_nombre(v):
    """Nombre (décimale française acceptée : « 100,00 »), None si vide/invalide."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(" ", "").replace("\u202f", "").replace("\xa0", "")
    if not s:
        return None
    s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _reporter(champ: str, valeur, data: dict) -> None:
    """Écrit la valeur dans data (première occurrence) si non replacée."""
    if champ not in data:
        data[champ] = valeur


# ---------------------------------------------------------------------------
def _feuille(wb, *mots) -> object:
    """Première feuille dont le nom normalisé contient tous les mots donnés."""
    for ws in wb.worksheets:
        n = normaliser(ws.title)
        if all(m in n for m in mots):
            return ws
    return None


def _lire_donnees_projet(ws, data: dict, warnings: list) -> None:
    for row in ws.iter_rows(values_only=True):
        libelle = row[0]
        if libelle is None or str(libelle).strip() == "":
            continue
        texte = str(libelle).strip()
        # En-têtes de section (« — 1. Identification — ») ignorées.
        if texte.startswith("—") or texte.startswith("-") or texte.startswith("*"):
            continue
        champ = _champ(libelle)
        if champ is None:
            continue
        valeur = row[1] if len(row) > 1 else None

        if champ in _CHAMPS_NUM:
            # « auto » / vide → on conserve la valeur par défaut (H_z auto).
            if valeur is None:
                continue
            if isinstance(valeur, str) and "auto" in valeur.lower():
                continue
            nb = _parse_nombre(valeur)
            if nb is None:
                warnings.append(f"« {texte} » : valeur numérique invalide (« {valeur} ») — ignorée.")
                continue
            _reporter(champ, nb, data)
        else:
            texte_valeur = "" if valeur is None else str(valeur).strip()
            if champ == "schema":
                texte_valeur = texte_valeur.upper().strip()
            if texte_valeur == "":
                continue
            _reporter(champ, texte_valeur, data)


def _lire_organes_client(ws, data: dict, warnings: list) -> None:
    organes = []
    for row in ws.iter_rows(values_only=True):
        typ = row[0]
        if typ is None or str(typ).strip() == "":
            continue
        n_typ = normaliser(typ)
        if not n_typ or n_typ.startswith("type"):
            continue  # en-tête ou ligne de note (« Type : trifon … »)
        cat = _TYPES_ALIASES.get(n_typ)
        if cat is None:
            warnings.append(f"Organes client : type inconnu « {typ} » — ligne ignorée.")
            continue
        dn = None
        if len(row) > 1:
            dn = _parse_nombre(row[1])
        if cat == "vanne":
            if dn is not None:
                data["dn_vanne_sectionnement"] = dn
                organes.append({"type": "vanne", "dn": dn, "nombre": 1,
                                "position": "", "fournisseur": ""})
            continue
        if dn is None:
            warnings.append(f"Organes client : DN manquant/invalide pour « {typ} » — ligne ignorée.")
            continue
        try:
            nb = int(float(row[2])) if (len(row) > 2 and row[2] not in (None, "")) else 1
        except (TypeError, ValueError):
            nb = 1
        nb = max(1, nb)
        pos = str(row[3]).strip().lower() if len(row) > 3 and row[3] is not None else ""
        four = str(row[4]).strip() if len(row) > 4 and row[4] is not None else ""
        entry = {"type": cat, "dn": int(dn), "nombre": nb, "position": pos}
        if four:
            entry["fournisseur"] = four
        organes.append(entry)
    if organes:
        data["organes_client"] = [
            o for o in organes if o.get("type") != "vanne"]
        vannes = [o for o in organes if o.get("type") == "vanne"]
        if vannes:
            data["dn_vanne_sectionnement"] = vannes[-1]["dn"]


# ---------------------------------------------------------------------------
def lire_donnees_projet(chemin: str):
    """Lit un classeur Excel et renvoie (dict_etat, avertissements).

    Le dict est directement compatible avec `EtatApplication.to_dict()` :
    seuls les champs renseignés dans le fichier sont présents.
    """
    warnings: list[str] = []
    data: dict = {}

    wb = openpyxl.load_workbook(chemin, data_only=True, read_only=True)
    try:
        ws_projet = _feuille(wb, "donnees", "projet") or wb.worksheets[0]
        _lire_donnees_projet(ws_projet, data, warnings)

        ws_organes = _feuille(wb, "organes", "client")
        if ws_organes is None:
            if wb.worksheets[0] is not ws_projet:
                ws_organes = wb.worksheets[0]
        if ws_organes is not None:
            _lire_organes_client(ws_organes, data, warnings)
    finally:
        wb.close()

    _controles_validite(data, warnings)
    return data, warnings


# ---------------------------------------------------------------------------
def _controles_validite(data: dict, warnings: list) -> None:
    schema = data.get("schema", "").upper()
    if not schema:
        warnings.append("Schéma de vidange non renseigné — le calcul n'est pas possible.")
    elif schema not in ALLOWED_SCHEMAS:
        warnings.append(f"Schéma de vidange « {schema} » inconnu (attendu : "
                        f"{', '.join(ALLOWED_SCHEMAS)}).")
    else:
        manquantes = [d for d in DISTANCES_PAR_CAS.get(schema, ())
                      if d not in data]
        if manquantes:
            warnings.append(
                f"Distances manquantes pour le schéma {schema} : "
                f"{', '.join(manquantes)}.")

    if "z_ve" not in data:
        warnings.append("Z_Ve (point haut) non renseigné.")
    if "z_vi1" not in data and "z_vi2" not in data:
        warnings.append("Ni Z_Vi1 ni Z_Vi2 renseignés — le profil est incomplet.")